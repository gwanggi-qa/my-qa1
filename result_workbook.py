"""Export run results into a copy of the uploaded OOXML package.

No device, AI, network, filesystem writes, or verdict inference belongs here.
Unedited ZIP members are copied byte-for-byte (including drawings and macros).
"""
from dataclasses import dataclass
import hashlib
import io
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape


MAX_WORKBOOK_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
STATUS_CODES = {"passed": "Pass", "failed": "F", "blocked": "NT",
                "stopped": "NT", "not_run": "NT", "not_applicable": "NA"}
ID_HEADERS = {"no", "번호", "tcid", "testcaseid", "tc번호", "통합id", "테스트id", "caseno"}
RESULT_HEADERS = {"검증결과", "검증결과pfntna", "result", "결과"}


def normalized(value):
    return re.sub(r"[\s_\-().,/]", "", str(value or "")).casefold()


@dataclass(frozen=True)
class RowTarget:
    index: int
    sheet: str
    row: int
    case_no: str
    id_cell: str
    result_cell: str


@dataclass(frozen=True)
class Template:
    payload: bytes
    file_hash: str
    filename: str
    rows: tuple
    formulas: tuple
    warnings: tuple


@dataclass(frozen=True)
class RunSnapshot:
    run_id: str
    template: Template
    statuses: tuple


def _open_package(payload):
    if not payload or len(payload) > MAX_WORKBOOK_BYTES:
        raise ValueError("결과 Excel은 20 MiB 이하 원본만 지원합니다.")
    package = zipfile.ZipFile(io.BytesIO(payload))
    infos = package.infolist()
    if len(infos) != len({i.filename for i in infos}) or sum(i.file_size for i in infos) > MAX_EXPANDED_BYTES:
        package.close()
        raise ValueError("중복 항목 또는 너무 큰 Excel 패키지는 지원하지 않습니다.")
    if any(i.filename.startswith("_xmlsignatures/") for i in infos):
        package.close()
        raise ValueError("전자서명된 Excel은 결과 사본 출력을 지원하지 않습니다.")
    return package


def _sheet_paths(package):
    root = ET.fromstring(package.read("xl/workbook.xml"))
    relations = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    lookup = {r.get("Id"): r.get("Target") for r in relations if r.get("TargetMode") != "External"}
    return {s.get("name"): posixpath.normpath("xl/" + lookup[s.get(f"{{{REL}}}id")])
            if not lookup[s.get(f"{{{REL}}}id")].startswith("/")
            else lookup[s.get(f"{{{REL}}}id")].lstrip("/")
            for s in root.find(f"{{{NS}}}sheets")}


def _headers(ws, selection=None):
    if selection and selection.get("sheet") == ws.title:
        mapping = selection.get("mapping", {})
        if "case_no" in mapping and "qa_result" in mapping:
            return selection["header_row"], mapping["case_no"] + 1, mapping["qa_result"] + 1
    for row in ws.iter_rows(max_row=min(ws.max_row, 50)):
        ids = [c.column for c in row if normalized(c.value) in ID_HEADERS]
        results = [c.column for c in row if normalized(c.value) in RESULT_HEADERS]
        if len(ids) == len(results) == 1:
            return next(c.row for c in row if getattr(c, "row", None)), ids[0], results[0]
    raise ValueError(f"{ws.title}: 원본 TC ID와 검증결과 열을 명확히 매핑해야 합니다.")


def _split_arguments(text):
    result, start, depth, quote = [], 0, 0, None
    i = 0
    while i < len(text):
        ch = text[i]
        if quote:
            if ch == quote:
                if i + 1 < len(text) and text[i + 1] == quote:
                    i += 1
                else:
                    quote = None
        elif ch in ('"', "'"):
            quote = ch
        elif ch in "({":
            depth += 1
        elif ch in ")}":
            depth -= 1
        elif ch == "," and depth == 0:
            result.append(text[start:i].strip())
            start = i + 1
        i += 1
    result.append(text[start:].strip())
    return result


def _count_calls(formula):
    # Skip quoted strings/sheet names; never rewrite text that merely looks like a formula.
    i, quote = 0, None
    while i < len(formula):
        ch = formula[i]
        if quote:
            if ch == quote:
                if i + 1 < len(formula) and formula[i + 1] == quote:
                    i += 1
                else:
                    quote = None
            i += 1
            continue
        if ch in ('"', "'"):
            quote = ch
            i += 1
            continue
        match = re.match(r"COUNTIFS?\s*\(", formula[i:], re.I)
        if match and (i == 0 or not (formula[i-1].isalnum() or formula[i-1] in "_.")):
            end, depth, inner_quote = i + match.end(), 1, None
            begin = end
            while end < len(formula) and depth:
                c = formula[end]
                if inner_quote:
                    if c == inner_quote:
                        if end + 1 < len(formula) and formula[end + 1] == c:
                            end += 1
                        else:
                            inner_quote = None
                elif c in ('"', "'"):
                    inner_quote = c
                elif c == "(":
                    depth += 1
                elif c == ")":
                    depth -= 1
                end += 1
            if depth == 0:
                yield i, end, match.group().split("(")[0].strip().upper(), _split_arguments(formula[begin:end-1])
                i = end
                continue
        i += 1


def _reference(text, sheet):
    match = re.fullmatch(r"(?:(?:'((?:[^']|'')+)'|([^!'\[\]]+))!)?(\$?[A-Za-z]{1,3}\$?\d*|\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3})", text.strip())
    if not match:
        return None
    return (match[1].replace("''", "'") if match[1] else match[2] or sheet,
            match[3].replace("$", "").upper())


def compatible_formula(formula, sheet, workbook, result_columns):
    """Only exact COUNTIF(S) success criteria on known result columns are changed."""
    calls, candidates = list(_count_calls(formula)), []
    for start, end, name, args in calls:
        if len(args) < 2 or len(args) % 2 or (name == "COUNTIF" and len(args) != 2):
            continue
        success = []
        for pos in range(0, len(args), 2):
            ref = _reference(args[pos], sheet)
            if not ref or ref[0] not in result_columns:
                continue
            cols = re.findall(r"[A-Z]+", ref[1])
            if not cols or len(set(cols)) != 1 or cols[0] not in result_columns[ref[0]]:
                continue
            criterion = args[pos+1]
            if criterion.startswith('"') and criterion.endswith('"'):
                value = criterion[1:-1].replace('""', '"')
            else:
                cell = _reference(criterion, sheet)
                value = workbook[cell[0]][cell[1]].value if cell and re.fullmatch(r"[A-Z]+\d+", cell[1]) and cell[0] in workbook.sheetnames else None
            if str(value or "").casefold() in {"p", "=p", "pass", "=pass"}:
                success.append(pos+1)
        if len(success) == 1:
            pos = success[0]
            legacy, current = list(args), list(args)
            legacy[pos], current[pos] = '"P"', '"Pass"'
            old_call = f"{name}({','.join(legacy)})"
            new_call = f"{name}({','.join(current)})"
            candidates.append((start, end, old_call, new_call))
    # Already compatible formulas must not be expanded a second time.
    canonical = {f"{name}({','.join(args)})".casefold() for _, _, name, args in calls}
    for start, end, old_call, new_call in reversed(candidates):
        if old_call.casefold() in canonical and new_call.casefold() in canonical:
            continue
        formula = formula[:start] + f"({old_call}+{new_call})" + formula[end:]
    return formula


def prepare_template(payload, filename, steps, selection=None):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    with _open_package(payload) as package:
        paths = _sheet_paths(package)
        merged = {}
        for sheet, path in paths.items():
            root = ET.fromstring(package.read(path))
            merged[sheet] = [item.get("ref") for item in root.findall(f"{{{NS}}}mergeCells/{{{NS}}}mergeCell")]
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
        try:
            headers, result_columns, targets, target_cells = {}, {}, [], {}
            for ws in workbook:
                try:
                    headers[ws.title] = _headers(ws, selection)
                    result_columns[ws.title] = {get_column_letter(headers[ws.title][2])}
                except ValueError:
                    pass
            for sheet in {step.get("sheet") for step in steps}:
                if sheet in headers:
                    needed = {step.get("row") for step in steps if step.get("sheet") == sheet}
                    for row_number, row_cells in enumerate(workbook[sheet].iter_rows(), 1):
                        if row_number in needed:
                            target_cells[(sheet, row_number)] = row_cells
            for index, step in enumerate(steps, 1):
                sheet, row, case_no = step.get("sheet"), step.get("row"), str(step.get("case_no") or "").strip()
                if sheet not in headers or type(row) is not int or not case_no:
                    raise ValueError("결과 Excel에는 원본 시트·행·TC ID·검증결과 열이 필요합니다.")
                header, id_column, result_column = headers[sheet]
                ws = workbook[sheet]
                id_cell, result_cell = f"{get_column_letter(id_column)}{row}", f"{get_column_letter(result_column)}{row}"
                cells = target_cells.get((sheet, row), ())
                if row <= header or len(cells) < max(id_column, result_column) or str(cells[id_column-1].value or "").strip() != case_no:
                    raise ValueError("원본 행과 TC ID가 일치하지 않습니다.")
                from openpyxl.worksheet.cell_range import CellRange
                if any(result_cell in CellRange(area) for area in merged[sheet]):
                    raise ValueError("병합된 검증결과 셀은 안전하게 기록할 수 없습니다.")
                if cells[result_column-1].data_type == "f":
                    raise ValueError("검증결과 셀에 수식이 있어 결과 기록을 중단합니다.")
                targets.append(RowTarget(index, sheet, row, case_no, id_cell, result_cell))
            if not targets or len({(t.sheet, t.result_cell) for t in targets}) != len(targets):
                raise ValueError("결과 대상 행이 없거나 중복되어 있습니다.")
            formulas, warnings = [], []
            for ws in workbook:
                # Read formulas from package XML to avoid expanding large formatted ranges.
                root = ET.fromstring(package.read(paths[ws.title]))
                for cell in root.findall(f".//{{{NS}}}sheetData/{{{NS}}}row/{{{NS}}}c"):
                    f = cell.find(f"{{{NS}}}f")
                    if f is None or not f.text:
                        continue
                    original = "=" + f.text
                    updated = compatible_formula(original, ws.title, workbook, result_columns)
                    if updated != original:
                        if f.attrib:
                            warnings.append(f"{ws.title}!{cell.get('r')}: 공유/배열 수식은 변경하지 않았습니다.")
                        else:
                            formulas.append((ws.title, cell.get("r"), updated[1:]))
            return Template(bytes(payload), hashlib.sha256(payload).hexdigest(), filename,
                            tuple(targets), tuple(formulas), tuple(warnings))
        finally:
            workbook.close()


def finish_run(run_id, template, steps, results, stopped=False, active_indices=(), selected_indices=None):
    if not re.fullmatch(r"[a-f0-9]{32}", run_id):
        raise ValueError("Invalid run ID")
    if len(steps) != len(template.rows):
        raise ValueError("실행 목록과 원본 행 수가 다릅니다.")
    for target, step in zip(template.rows, steps):
        if (step.get("sheet"), step.get("row"), str(step.get("case_no") or "").strip()) != (target.sheet, target.row, target.case_no):
            raise ValueError("실행 중 원본 행 식별자가 변경되었습니다.")
    selected = set(range(1, len(steps)+1)) if selected_indices is None else set(selected_indices)
    if not selected or any(type(i) is not int or not 1 <= i <= len(steps) for i in selected):
        raise ValueError("선택한 원본 TC 인덱스가 올바르지 않습니다.")
    statuses = ["preserved" if i not in selected else "stopped" if stopped and i in active_indices else "not_run" for i in range(1, len(steps)+1)]
    seen = set()
    for result in results:
        index, status = result.get("index"), result.get("status")
        if type(index) is not int or index not in selected or index in seen or status not in STATUS_CODES:
            raise ValueError("중복되거나 유효하지 않은 실행 결과입니다.")
        seen.add(index)
        statuses[index-1] = status
    return RunSnapshot(run_id, template, tuple(statuses))


CELL_RE = re.compile(r'<(?P<tag>(?:[\w]+:)?c)\b(?P<attrs>[^>]*?)(?:\s*/>|>(?P<body>.*?)</(?P=tag)\s*>)', re.S)


def _cell_match(xml, address):
    for match in CELL_RE.finditer(xml):
        attr = re.search(r'\br\s*=\s*["\']([^"\']+)["\']', match['attrs'])
        if attr and attr[1] == address:
            return match
    return None


def _write_value(xml, address, value):
    match = _cell_match(xml, address)
    if match:
        tag, attrs = match['tag'], re.sub(r'\s+t\s*=\s*["\'][^"\']*["\']', '', match['attrs'])
        prefix = tag[:-1]
        replacement = f'<{tag}{attrs} t="inlineStr"><{prefix}is><{prefix}t>{escape(value)}</{prefix}t></{prefix}is></{tag}>'
        return xml[:match.start()] + replacement + xml[match.end():]
    row_number = re.search(r"\d+$", address)[0]
    row_re = re.compile(r'<(?P<tag>(?:[\w]+:)?row)\b(?P<attrs>[^>]*)>(?P<body>.*?)</(?P=tag)\s*>', re.S)
    from openpyxl.utils import column_index_from_string
    col = column_index_from_string(re.match(r"[A-Z]+", address)[0])
    for row in row_re.finditer(xml):
        attr = re.search(r'\br\s*=\s*["\']([^"\']+)["\']', row['attrs'])
        if not attr or attr[1] != row_number:
            continue
        prefix = row['tag'][:-3]
        new_cell = f'<{prefix}c r="{address}" t="inlineStr"><{prefix}is><{prefix}t>{escape(value)}</{prefix}t></{prefix}is></{prefix}c>'
        offset = row.end('body')
        for cell in CELL_RE.finditer(row['body']):
            cell_ref = re.search(r'\br\s*=\s*["\']([A-Z]+)\d+["\']', cell['attrs'])
            if cell_ref and column_index_from_string(cell_ref[1]) > col:
                offset = row.start('body') + cell.start()
                break
        return xml[:offset] + new_cell + xml[offset:]
    raise ValueError("원본 결과 행을 찾을 수 없습니다.")


def _write_formula(xml, address, formula):
    match = _cell_match(xml, address)
    if not match:
        raise ValueError("집계 수식 셀을 찾을 수 없습니다.")
    cell = match.group()
    cell, count = re.subn(r'<((?:[\w]+:)?f)>(.*?)</\1\s*>', lambda m: f'<{m[1]}>{escape(formula)}</{m[1]}>', cell, count=1, flags=re.S)
    if count != 1:
        raise ValueError("지원되지 않는 집계 수식 형식입니다.")
    cell = re.sub(r'<((?:[\w]+:)?v)\b[^>]*(?:/>|>.*?</\1\s*>)', '', cell, flags=re.S)
    return xml[:match.start()] + cell + xml[match.end():]


def export_run(snapshot, expected_hash):
    template = snapshot.template
    if expected_hash != template.file_hash or hashlib.sha256(template.payload).hexdigest() != expected_hash:
        raise ValueError("원본 파일 해시가 일치하지 않습니다.")
    if len(snapshot.statuses) != len(template.rows) or any(s not in {*STATUS_CODES, "preserved"} for s in snapshot.statuses):
        raise ValueError("결과 스냅샷이 유효하지 않습니다.")
    with _open_package(template.payload) as source:
        paths = _sheet_paths(source)
        edits = {}
        for target, status in zip(template.rows, snapshot.statuses):
            if status == "preserved":
                continue
            path = paths[target.sheet]
            xml = edits.get(path, source.read(path).decode("utf-8"))
            edits[path] = _write_value(xml, target.result_cell, STATUS_CODES[status])
        for sheet, cell, formula in template.formulas:
            path = paths[sheet]
            xml = edits.get(path, source.read(path).decode("utf-8"))
            edits[path] = _write_formula(xml, cell, formula)
        workbook = source.read("xl/workbook.xml").decode("utf-8")
        calc = re.search(r'<((?:[\w]+:)?calcPr)\b([^>]*?)(/>)', workbook)
        if calc:
            attrs = re.sub(r'\s+(?:calcMode|fullCalcOnLoad|forceFullCalc)="[^"]*"', '', calc[2])
            replacement = f'<{calc[1]}{attrs} calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
            workbook = workbook[:calc.start()] + replacement + workbook[calc.end():]
        else:
            # Insert at the schema position before optional trailing elements.
            close = re.search(r'</((?:[\w]+:)?workbook)>', workbook)
            prefix = close[1][:-8]
            insertion = re.search(r'<(?:[\w]+:)?(?:oleSize|customWorkbookViews|pivotCaches|smartTagPr|smartTagTypes|webPublishing|fileRecoveryPr|webPublishObjects|extLst)\b', workbook)
            at = insertion.start() if insertion else close.start()
            workbook = workbook[:at] + f'<{prefix}calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>' + workbook[at:]
        edits["xl/workbook.xml"] = workbook
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as dest:
            for info in source.infolist():
                content = edits[info.filename].encode("utf-8") if info.filename in edits else source.read(info.filename)
                dest.writestr(info, content)
        return output.getvalue()

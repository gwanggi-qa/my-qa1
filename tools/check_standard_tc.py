"""Generate a synthetic standard-format result for offline visual inspection."""
import hashlib
import io
from pathlib import Path
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from openpyxl import load_workbook
import narrative_excel
import result_workbook
import tc_format

if __name__=='__main__':
    if len(sys.argv)>1:
        original=Path(sys.argv[1])
        data=original.read_bytes()
        before=hashlib.sha256(data).hexdigest()
        book=load_workbook(io.BytesIO(data),data_only=False)
        assert tuple(c.value for c in book['TestCase'][1])==tc_format.HEADERS
        selection=narrative_excel.detect(book)
        steps=narrative_excel.parse(book,selection)
        assert len(steps)==10
        assert steps[0]['qa_result']=='PASS'
        assert all(s.get('automation_result','')=='' for s in steps)
        book.close()
        assert hashlib.sha256(original.read_bytes()).hexdigest()==before
        print('Source: 14 standard headers, 10 rows, manual result retained, original unchanged.')
    statuses=['passed','failed','blocked']
    notes=['LC/RC 모두 확인됨 (모의 결과)', 'RC 영역이 없음 (모의 결과)', '검증 데이터 없음 (모의 결과)']
    steps=[{'case_no':f'DEMO-{i}','sheet':'TestCase','row':i+1,'title':'합성 TC: 학습 결과 표시',
            'priority':'P1','depth1':'학습 관리','depth2':'실력 분석','depth3':'LC/RC 카드','check_type':'UI-UX',
            'test_purpose':'LC와 RC 결과가 함께 표시되는지 확인','precondition':'합성 학습 관리 화면 진입',
            'test_item':'1. 학습 관리 화면을 연다.\n2. LC/RC 결과 영역을 확인한다.',
            'expected':'LC와 RC 결과 영역이 모두 표시된다.','qa_result':'수동 기록 유지',
            'note':'시각 검증용 합성 데이터. 실제 기기에서 실행한 결과가 아님.'} for i in range(1,4)]
    book=tc_format.workbook(tc_format.step_row(s) for s in steps)
    book.create_sheet('Sources').append(['Synthetic data only; no company reference.'])
    raw=tc_format.save(book)
    template=result_workbook.prepare_template(raw,'synthetic.xlsx',steps)
    snapshot=result_workbook.finish_run('c'*32,template,steps,[{'index':i,'status':s} for i,s in enumerate(statuses,1)])
    run={'run_id':'c'*32,'results':[{'source':{'sheet':s['sheet'],'row':s['row'],'tc_id':s['case_no'],'expected':s['expected']},
         'status':status,'actual':note,'reason':'','evidence':[],'stages':[{'number':1,'source':s['test_item'],
         'status':status,'actual':note,'evidence':[]}]} for s,status,note in zip(steps,statuses,notes)]}
    output=tc_format.append_execution_details(result_workbook.export_run(snapshot,template.file_hash),run)
    folder=ROOT/'qa_results'/'standard-tc-validation'
    folder.mkdir(exist_ok=True,parents=True)
    (folder/'standard-results.xlsx').write_bytes(output)
    print('Synthetic result workbook:',folder/'standard-results.xlsx')

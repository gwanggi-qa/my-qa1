"""Safety regressions: synthetic files/threads only; outer runner denies device/network."""
import contextlib
import io
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from unittest.mock import Mock, patch
import openpyxl
import mirror_server as m
import workbook_safety as safe
import execution_results as records
from test_result_workbook import fixture

class UploadSafetyTests(unittest.TestCase):
    def package(self, sheet):
        payload=io.BytesIO()
        with zipfile.ZipFile(payload,'w',zipfile.ZIP_DEFLATED) as z:
            z.writestr('[Content_Types].xml','<Types/>');z.writestr('xl/workbook.xml','<workbook/>')
            z.writestr('xl/worksheets/sheet1.xml',sheet)
        return payload.getvalue()
    def test_supported_original_workbook_unchanged(self):
        payload,_=fixture();source=io.BytesIO(payload)
        self.assertEqual(safe.read_upload(source,'원본.xlsx'),payload)
        self.assertEqual(source.tell(),0)
    def test_read_is_bounded_before_xml_parser(self):
        stream=Mock();stream.read.return_value=b'x'*129
        with patch.object(safe,'MAX_WORKBOOK_BYTES',128):
            with self.assertRaises(ValueError):safe.read_upload(stream,'bad.xlsx')
        stream.read.assert_called_once_with(129)
    def test_filename_paths_are_rejected(self):
        for name in ['../source.xlsx','C:\\source.xlsx','folder/source.xlsx','source.xlsx:ads','bad\x00.xlsx']:
            with self.subTest(name=name),self.assertRaises(ValueError):safe.read_upload(io.BytesIO(b'x'),name)
    def test_expanded_package_rejected_before_load_workbook(self):
        payload=self.package('<worksheet>'+(' '*1000)+'</worksheet>')
        with patch.object(safe,'MAX_EXPANDED_BYTES',100),patch('openpyxl.load_workbook') as parser:
            with self.assertRaises(ValueError):m.parse_excel(type('File',(),{'stream':io.BytesIO(payload),'filename':'limit.xlsx'})())
            parser.assert_not_called()
    def test_sparse_huge_dimensions_and_dtd_rejected(self):
        for xml in ['<worksheet><dimension ref="A1:XFD1048576"/></worksheet>',
                    '<worksheet><sheetData><row r="999999"><c r="A999999"/></row></sheetData></worksheet>',
                    '<!DOCTYPE worksheet [<!ENTITY unsafe "x">]><worksheet/>']:
            with self.subTest(xml=xml),self.assertRaises(ValueError):safe.validate_workbook(self.package(xml))
    def test_zip_path_and_duplicate_entries_rejected(self):
        for names in [['../outside.xml'],['xl/worksheets/../../x.xml'],['xl/worksheets/sheet1.xml']]:
            payload=io.BytesIO(self.package('<worksheet/>'))
            with zipfile.ZipFile(payload,'a') as z:
                for name in names:z.writestr(name,'<x/>')
            with self.assertRaises(ValueError):safe.validate_workbook(payload.getvalue())
    def test_preview_and_upload_reject_limits_before_parse_and_preserve_state(self):
        payload=self.package('<worksheet>'+(' '*1000)+'</worksheet>');old=object()
        for preview in ['true','false']:
            with patch.object(safe,'MAX_EXPANDED_BYTES',100),patch.object(m,'parse_excel') as parser,patch('openpyxl.load_workbook') as loader,patch.object(m,'_verified_run',old):
                response=m.app.test_client().post('/api/testcase',data={'file':(io.BytesIO(payload),'large.xlsx'),'preview':preview})
                self.assertEqual(response.status_code,400);self.assertIs(m._verified_run,old)
                parser.assert_not_called();loader.assert_not_called()
    def test_request_body_limit_before_upload_processing(self):
        with patch.object(safe,'MAX_WORKBOOK_BYTES',8),patch.object(m,'parse_excel') as parser:
            response=m.app.test_client().post('/api/testcase',data=b'x'*(1024*1024+9),content_type='multipart/form-data; boundary=fixture')
            self.assertEqual(response.status_code,413);parser.assert_not_called()
    def test_artifacts_are_confined_and_symlinks_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            for name in ['../x.png','a/b.png','a\\b.png','x.png:ads']:
                with self.subTest(name=name),self.assertRaises(ValueError):safe.artifact_path(folder,name)
            with patch.object(Path,'is_symlink',return_value=True):
                with self.assertRaises(ValueError):safe.artifact_path(folder,'x.png')
    def test_screenshot_endpoint_rejects_escape_without_reading_target(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(m,'QA_RESULTS_DIR',Path(folder)):
            (Path(folder)/'valid.png').write_bytes(b'fixture')
            client=m.app.test_client()
            self.assertEqual(client.get('/api/results/screenshot/valid.png').status_code,200)
            for name in ['../outside.png','a%5Cb.png','valid.png:ads']:
                self.assertIn(client.get('/api/results/screenshot/'+name).status_code,[400,404])
            with patch.object(safe,'artifact_path',side_effect=ValueError('escape')):
                self.assertEqual(client.get('/api/results/screenshot/valid.png').status_code,400)
    def test_new_recording_excel_strings_are_data_not_formulas(self):
        with tempfile.TemporaryDirectory() as folder,patch.object(m,'RECORDINGS_DIR',Path(folder)):
            data={'id':'recording-20261008-120000','actions':[{'type':'=1+1','x':'@SUM(A1:A2)','y':3}]}
            path=m.export_recording_excel(data);wb=openpyxl.load_workbook(path,data_only=False)
            self.assertEqual(wb.active['B2'].value,'=1+1');self.assertEqual(wb.active['B2'].data_type,'s')
            self.assertEqual(wb.active['C2'].data_type,'s');self.assertEqual(wb.active['D2'].value,3);wb.close()

class CancelOnEnter:
    def __init__(self,event):self.event=event
    def __enter__(self):self.event.set()
    def __exit__(self,*args):pass

class InputSafetyTests(unittest.TestCase):
    def setUp(self):
        self.stack=contextlib.ExitStack();self.addCleanup(self.stack.close)
        for name in ['_test_running','_random_running','_record_play_running']:self.stack.enter_context(patch.object(m,name,False))
        for name in ['_test_stop_event','_random_stop_event','_record_play_stop_event']:self.stack.enter_context(patch.object(m,name,threading.Event()))
        for name in ['_automation_start_lock','_adb_input_lock']:self.stack.enter_context(patch.object(m,name,threading.Lock()))
        self.stack.enter_context(patch.object(m,'_reviewed_plan',None));self.stack.enter_context(patch.object(m,'check_device',return_value=True))
        self.adb=self.stack.enter_context(patch.object(m,'adb',return_value=(b'',b'',0)))
        self.payload={'type':'tap','x':20,'y':20,'screen_width':100,'screen_height':100}
    def test_manual_input_cannot_interleave_any_automation(self):
        for flag in ['_test_running','_random_running','_record_play_running']:
            with patch.object(m,flag,True):
                self.assertEqual(m.app.test_client().post('/api/input',json=self.payload).status_code,409)
        self.adb.assert_not_called()
    def test_recorder_cancel_at_dispatch_lock_sends_no_input(self):
        m._record_play_running=True
        with patch.object(m,'_adb_input_lock',CancelOnEnter(m._record_play_stop_event)):
            with self.assertRaises(m.RecordingStopped):m.run_manual_input(self.payload,owner='recorder')
        self.adb.assert_not_called()
    def test_random_cancel_at_dispatch_lock_sends_no_input(self):
        m._random_running=True
        with patch.object(m,'capture_screen_hash',return_value='hash'),patch.object(m,'_adb_input_lock',CancelOnEnter(m._random_stop_event)):
            with self.assertRaises(m.RandomStopped):m.execute_random_action({'step':{'action':'tap','x':20,'y':20,'target':'fixture'}})
        self.adb.assert_not_called()
    def test_stop_acknowledgement_waits_for_inflight_input(self):
        for endpoint,event in [('/api/random/stop',m._random_stop_event),('/api/recordings/play/stop',m._record_play_stop_event)]:
            entered=threading.Event();release=threading.Event();done=threading.Event();responses=[]
            def input_work():
                with m._adb_input_lock:entered.set();release.wait(2)
            def stop_work():
                responses.append(m.app.test_client().post(endpoint).get_json());done.set()
            worker=threading.Thread(target=input_work);stopper=threading.Thread(target=stop_work)
            worker.start();self.assertTrue(entered.wait(1));stopper.start()
            try:self.assertFalse(done.wait(.05))
            finally:release.set();worker.join(2);stopper.join(2)
            self.assertTrue(done.is_set());self.assertTrue(event.is_set());self.assertTrue(responses[0]['input_quiescent'])
        self.adb.assert_not_called()
    def test_api_preserves_finished_review_stop_state(self):
        from types import SimpleNamespace
        plan=SimpleNamespace(approved=True,finished=True,stopped=True,run_id='fixture-stop',
            template=SimpleNamespace(file_hash='hash',filename='fixture.xlsx'),steps=[],results=[])
        with patch.object(m,'_verified_run',None),patch.object(m,'_reviewed_plan',plan),patch.object(records,'read_history',return_value=([],[])):
            response=m.app.test_client().get('/api/execution/results').get_json()
            self.assertEqual(response['current']['state'],'stopped')
            self.assertEqual(response['current']['method'],'manual_observation')

    def test_stopped_reviewed_history_remains_stopped(self):
        value=records.normalize({'run_id':'fixture','stopped':True,'verification_method':'operator-reviewed-local','results':[]})
        self.assertEqual(value['state'],'stopped');self.assertEqual(value['method'],'manual_observation')

if __name__=='__main__':unittest.main()
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook
from PIL import Image

import app_analysis as a
import reference_materials as r
from local_access import LocalAccessPolicy, HEADER
from test_app_analysis import apk


class ReferenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = a.AnalysisStore(Path(self.temp.name))
        self.build, _ = self.store.register(apk(), 'a.apk')
        self.sha = self.build['sha256']
        self.refs = r.ReferenceStore(self.store)

    def add(self, data=b'# Requirements\nExpected result', name='plan.md', kind='requirements', url=''):
        return self.refs.add(self.sha, io.BytesIO(data), name, kind, url)

    def test_text_dedup_provenance_and_delete(self):
        first, duplicate = self.add()
        self.assertFalse(duplicate)
        self.assertEqual(first['preview']['units'][0]['lines'][1]['line'], 2)
        second, duplicate = self.add()
        self.assertTrue(duplicate)
        self.assertEqual(first['id'], second['id'])
        self.add(url='https://www.notion.so/team/spec')
        self.add(kind='design')
        self.assertEqual(len(self.refs.list(self.sha)), 3)
        self.refs.delete(self.sha, first['id'])
        self.assertEqual(len(self.refs.list(self.sha)), 2)
        with self.assertRaises(a.AnalysisError):
            self.refs.get(self.sha, first['id'])

    def test_reference_scope_and_file_version(self):
        first, _ = self.add()
        changed, _ = self.add(data=b'new version')
        self.assertNotEqual(first['sha256'], changed['sha256'])
        other, _ = self.store.register(apk({'AndroidManifest.xml': b'other'}), 'other.apk')
        with self.assertRaises(a.AnalysisError) as caught:
            self.refs.get(other['sha256'], first['id'])
        self.assertEqual(caught.exception.status, 404)
        self.assertEqual(self.refs.list(other['sha256']), [])

    def test_bad_metadata_and_content_leave_existing_intact(self):
        first, _ = self.add()
        for args in [dict(name='a.html'), dict(kind='invalid'), dict(data=b'\x00binary'),
                     dict(data=b'\xff\xfe'), dict(data=b'', name='empty.md'),
                     dict(url='http://localhost:5000/'), dict(url='https://figma.com.evil.example/file'),
                     dict(url='https://secret@notion.so/page'), dict(url='file:///C:/secret')]:
            with self.subTest(args=args), self.assertRaises(a.AnalysisError):
                self.add(**args)
        self.assertEqual(len(self.refs.list(self.sha)), 1)
        self.assertEqual(self.refs.get(self.sha, first['id'])['filename'], 'plan.md')

    def test_file_and_storage_limits_and_parser_busy(self):
        with patch.object(r, 'MAX_FILE', 5), self.assertRaises(a.AnalysisError) as caught:
            self.add()
        self.assertEqual(caught.exception.status, 413)
        with patch.object(r, 'MAX_TOTAL', 1), self.assertRaises(a.AnalysisError):
            self.add()
        self.refs.lock.acquire()
        try:
            with self.assertRaises(a.AnalysisError) as caught:
                self.add()
            self.assertEqual(caught.exception.status, 409)
        finally:
            self.refs.lock.release()
        self.assertEqual(self.refs.list(self.sha), [])

    def test_text_truncation_and_markup_are_explicit(self):
        with patch.object(r, 'MAX_TEXT', 10):
            value = r.text_preview(b'<script>alert(1)</script>')
        self.assertEqual(value['status'], 'partial')
        self.assertTrue(value['warnings'])
        self.assertTrue(value['units'][0]['lines'][0]['text'].startswith('<script>'))

    def test_workbook_preserves_coordinates_formulas_merges_hidden_state(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'TC'
        sheet['A1'] = 'TC ID'
        sheet['B1'] = 'Expected'
        sheet['A2'] = 'TC-001'
        sheet['B2'] = 'Line 1\nLine 2'
        sheet.merge_cells('A3:B3')
        sheet['A3'] = 'merged'
        sheet['C2'] = '=1+1'
        hidden = book.create_sheet('Hidden')
        hidden.sheet_state = 'hidden'
        hidden['A1'] = 'hidden requirements'
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        item, _ = self.add(stream.getvalue(), 'cases.xlsx', 'testcase')
        units = item['preview']['units']
        self.assertEqual(units[0]['merged_ranges'], ['A3:B3'])
        self.assertEqual(units[1]['sheet_state'], 'hidden')
        cells = {c['cell']: c for row in units[0]['rows'] for c in row['cells']}
        self.assertEqual(cells['B2']['value'], 'Line 1\nLine 2')
        self.assertTrue(cells['C2']['formula'])
        self.assertEqual(cells['C2']['value'], '=1+1')
        with self.assertRaises(a.AnalysisError):
            self.add(b'not xlsx', 'bad.xlsx')

    def test_image_normalized_and_unsupported_rejected(self):
        stream = io.BytesIO()
        Image.new('RGB', (100, 50), 'red').save(stream, 'PNG')
        item, _ = self.add(stream.getvalue(), 'screen.png', 'design')
        self.assertEqual(item['preview']['status'], 'visual_only')
        image = self.refs.get(self.sha, item['id'], image=True)
        self.assertTrue(image.startswith(b'\x89PNG'))
        with self.assertRaises(a.AnalysisError):
            self.add(b'not image', 'bad.jpg')
        with self.assertRaises(a.AnalysisError):
            self.refs.get(self.sha, 'invalid')

    def test_pdf_worker_contract_timeout_and_missing_dependency(self):
        class Completed:
            returncode = 0
            stdout = json.dumps({'format': 'pdf', 'status': 'ready', 'units': []}).encode()
        with patch('importlib.util.find_spec', return_value=True), patch.object(r.subprocess, 'run', return_value=Completed()) as run:
            value = r.pdf_preview(b'%PDF-fixture')
            self.assertEqual(value['format'], 'pdf')
            self.assertEqual(run.call_args.kwargs['timeout'], 30)
            self.assertNotIn('OPENAI_API_KEY', run.call_args.kwargs['env'])
        with patch('importlib.util.find_spec', return_value=True), patch.object(r.subprocess, 'run', side_effect=subprocess.TimeoutExpired('pdf', 30)):
            with self.assertRaises(a.AnalysisError) as caught:
                r.pdf_preview(b'%PDF-fixture')
            self.assertEqual(caught.exception.status, 422)
        with patch('importlib.util.find_spec', return_value=None):
            with self.assertRaises(a.AnalysisError) as caught:
                r.pdf_preview(b'%PDF-fixture')
            self.assertEqual(caught.exception.status, 503)

    def test_pdf_scanned_and_invalid_worker_response(self):
        class Page:
            def extract_text(self):
                return ''
        class Reader:
            is_encrypted = False
            pages = [Page()]
        with patch('pypdf.PdfReader', return_value=Reader()):
            value = r.extract_pdf('unused')
        self.assertEqual(value['status'], 'partial')
        self.assertEqual(value['units'][0]['page'], 1)
        class Completed:
            returncode = 0
            stdout = b'invalid json'
        with patch('importlib.util.find_spec', return_value=True), patch.object(r.subprocess, 'run', return_value=Completed()):
            with self.assertRaises(a.AnalysisError):
                r.pdf_preview(b'%PDF-fixture')

    def test_schema_one_upgrade_preserves_build(self):
        with self.store.connect() as db:
            db.execute('DROP TABLE reference_materials')
            db.execute('PRAGMA user_version=1')
        upgraded = a.AnalysisStore(self.store.root)
        self.assertEqual(upgraded.get(self.sha)['filename'], 'a.apk')
        self.assertEqual(r.ReferenceStore(upgraded).list(self.sha), [])
        with upgraded.connect() as db:
            self.assertEqual(db.execute('PRAGMA user_version').fetchone()[0], 3)


class ReferenceApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.app = Flask('references-test')
        self.policy = LocalAccessPolicy('127.0.0.1', 5000, ['localhost', '127.0.0.1'])
        self.policy.install(self.app)
        a.register(self.app, self.policy, self.temp.name)
        self.client = self.app.test_client()
        store = self.app.extensions['app_analysis'].store
        self.sha = store.register(apk(), 'a.apk')[0]['sha256']
        self.url = '/api/app-analysis/builds/' + self.sha + '/references'
        self.headers = {HEADER: self.policy.token, 'Origin': 'http://localhost:5000'}

    def call(self, url=None, **kwargs):
        kwargs.setdefault('headers', self.headers)
        return self.client.open(url or self.url, base_url='http://localhost:5000', **kwargs)

    def test_full_lifecycle_and_header_guards(self):
        value = self.call(method='POST', data={'file': (io.BytesIO(b'Expected'), 'req.md'), 'kind': 'requirements'})
        self.assertEqual(value.status_code, 201)
        ref_id = value.json['reference']['id']
        self.assertEqual(len(self.call().json['references']), 1)
        url = self.url + '/' + ref_id
        self.assertEqual(self.call(url).status_code, 200)
        for headers in [{}, {**self.headers, 'Origin': 'http://evil.example'}, {**self.headers, 'X-Forwarded-For': '127.0.0.1'}]:
            self.assertEqual(self.call(url, method='DELETE', headers=headers, json={}).status_code, 403)
        self.assertEqual(self.call(url, method='DELETE', json={'other': True}).status_code, 400)
        self.assertEqual(self.call(url, method='DELETE', json={}).status_code, 200)
        self.assertEqual(self.call(url).status_code, 404)

    def test_multipart_contract(self):
        self.assertEqual(self.call(method='POST', json={}).status_code, 415)
        for data in [{'kind': 'requirements'}, {'file': (io.BytesIO(b'x'), 'x.txt'), 'kind': 'invalid'},
                     {'file': (io.BytesIO(b'x'), 'x.txt'), 'kind': 'requirements', 'path': 'C:/secret'}]:
            response = self.call(method='POST', data=data, content_type='multipart/form-data')
            self.assertEqual(response.status_code, 400)


if __name__ == '__main__':
    unittest.main()

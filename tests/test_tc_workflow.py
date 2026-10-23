import copy
import io
import tempfile
import threading
import zipfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
from flask import Flask
from openpyxl import load_workbook, Workbook
from werkzeug.datastructures import FileStorage
from PIL import Image
import mirror_server as m
import tc_workflow as w
import execution_results as records
import tc_format
import result_workbook
from local_access import LocalAccessPolicy, HEADER


def step():
    return {'source':'narrative_qa','case_no':'DTC-001','sheet':'TestCase','row':2,'action':'blocked',
            'ai_pending':True,'title':'설정 표시','precondition':'"홈"이 표시된다',
            'test_item':'1. 화면을 확인한다.','expected':'"설정"이 표시된다','qa_expected':'"설정"이 표시된다'}


def action(kind='inspect',target='',procedure=1):
    return {'procedure':procedure,'action':kind,'target':target,'direction':'','wait_ms':0}


def ai_plan(s,actions=None,index=1):
    actions=actions or [action()]
    return {'index':index,'actions':[{**a,'basis':'procedure','source':w.procedures(s)[a['procedure']-1]} for a in actions],
            'checks':[{'expectation':i,'after_action':len(actions)} for i,_ in enumerate(w.tc_adaptive.expectations(s),1)],
            'blocked_reason':''}


class WorkflowTests(unittest.TestCase):
    def test_back_guard_distinguishes_document_prose_from_confirmation(self):
        import xml.etree.ElementTree as ET
        prose='개인정보의 저장 및 삭제에 대한 설명입니다. ' * 8
        document='<hierarchy><node class="android.webkit.WebView"><node text="'+prose+'"/></node></hierarchy>'
        self.assertFalse(w.tc_adaptive.back_confirmation_risk(ET.fromstring(document)))
        self.assertFalse(w.tc_adaptive.back_confirmation_risk(ET.fromstring(
            document.replace('</node></hierarchy>','<node text="결제정보"/></node></hierarchy>'))))
        for xml in [document.replace('<node text=', '<node clickable="true" text='),
                    '<hierarchy><node text="삭제하시겠습니까?"/></hierarchy>',
                    '<hierarchy><node clickable="true"><node text="저장"/></node></hierarchy>']:
            self.assertTrue(w.tc_adaptive.back_confirmation_risk(ET.fromstring(xml)))

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        png=io.BytesIO()
        Image.new('RGB',(100,100),'white').save(png,'PNG')
        self.metadata={'upload_id':'v1','file_hash':'hash1','filename':'generated.xlsx'}
        xml='<hierarchy><node package="fixture" text="홈"/><node text="설정" clickable="true" enabled="true" bounds="[1,1][20,20]"/></hierarchy>'
        self.host=SimpleNamespace(_test_metadata=self.metadata,_test_steps=[step()],_test_workbook=None,
            _test_lock=threading.RLock(),_automation_start_lock=threading.RLock(),_adb_input_lock=threading.Lock(),
            _test_stop_event=threading.Event(),_test_ai_context=threading.local(),_test_running=False,
            _random_running=False,_record_play_running=False,_verified_run=None,
            socketio=Mock(),QA_RESULTS_DIR=Path(self.temp.name),check_device=Mock(return_value=True),
            get_device_info=Mock(return_value={'android':'16'}),foreground_app_evidence=Mock(return_value={'package':'fixture'}),
            capture_screen_png=Mock(return_value=png.getvalue()),dump_ui_xml=Mock(return_value=xml),
            qa_result_screenshot_url=lambda p:'/api/results/screenshot/'+p.name,
            emit_progress=Mock(),StepResult=m.StepResult,BlockedStep=m.BlockedStep,append_test_run_history=Mock(),
            request_test_ai_json=Mock(),verify_visual_step=Mock(),adb=Mock(return_value=(b'',b'',0)),
            require_review_reset_confirmation=Mock(return_value='fixture-token'),begin_redmine_review_run=Mock(),
            record_tc_review_result=Mock(),ReviewResetRequired=m.ReviewResetRequired,
            ai_candidate_is_safe=m.ai_candidate_is_safe,text_value_matches=m.text_value_matches,
            parse_bounds=m.parse_bounds,bounds_center=m.bounds_center)
        candidate={'candidate_id':'c001','label':'설정','action':'tap','bounds':'[1,1][20,20]','source':'ui-analysis'}
        self.host.ai_test_screen_candidates=Mock(return_value=([candidate],[]))
        self.app=Flask(__name__)
        self.policy=LocalAccessPolicy('127.0.0.1',5000,['localhost'])
        self.policy.install(self.app)
        self.service=w.register(self.app,lambda:self.host)
        self.client=self.app.test_client()
        self.headers={HEADER:self.policy.token,'Origin':'http://localhost:5000'}

    def call(self,path,body=None):
        return self.client.open('/api/tc-workflow/'+path,method='GET' if body is None else 'POST',
                                json=body,base_url='http://localhost:5000',headers=self.headers)

    def prepare(self,ai=False):
        response=self.call('prepare',{'upload_id':self.host._test_metadata['upload_id'],'file_hash':self.host._test_metadata['file_hash'],'indices':[1],
                                      'ai_mode':ai,'consent':ai})
        self.assertEqual(response.status_code,202,response.json)
        flow=self.service['current']
        flow.worker.join(3)
        self.assertFalse(flow.worker.is_alive())
        return flow

    def start(self,flow):
        response=self.call(flow.id+'/start',{'plan_digest':flow.digest,'approved':True})
        self.assertEqual(response.status_code,202,response.json)
        flow.worker.join(3)
        self.assertFalse(flow.worker.is_alive())
        return flow.snapshot()

    def visual(self,probe):
        probe['_visual_evidence']={'screenshot_url':'/fixture.png','screenshot_file':'fixture.png',
                                    'findings':[{'id':'v1','status':'met'}]}
        return '모든 원문 조건 확인'

    def test_expected_without_result_column_and_export(self):
        flow=self.prepare()
        self.assertEqual(flow.state,'ready')
        self.assertEqual(flow.plans[0]['blocked_reason'],'')
        self.host.adb.assert_not_called()
        result=self.start(flow)
        self.assertEqual(result['counts']['passed'],1)
        self.assertEqual(result['results'][0]['source']['expected'],step()['expected'])
        exported=self.call(flow.id+'/export')
        book=load_workbook(io.BytesIO(exported.data))
        self.assertEqual(book.active['K2'].value,step()['expected'])
        self.assertEqual(book.active['M2'].value,'PASS')
        book.close()
        self.host.request_test_ai_json.assert_not_called()
        self.host.begin_redmine_review_run.assert_called_once_with(flow.data['run_id'],'fixture-token')

    def test_review_confirmation_and_device_failure_preserve_queue(self):
        flow=self.prepare()
        self.host.require_review_reset_confirmation.side_effect=m.ReviewResetRequired({'pending':1,'token':'version'})
        response=self.call(flow.id+'/start',{'plan_digest':flow.digest,'approved':True})
        self.assertEqual(response.status_code,409)
        self.assertEqual(response.json['code'],'fail_review_reset_required')
        self.assertEqual(flow.state,'ready')
        self.host.begin_redmine_review_run.assert_not_called()
        self.host.require_review_reset_confirmation.side_effect=None
        self.host.check_device.return_value=False
        self.start(flow)
        self.host.begin_redmine_review_run.assert_not_called()

    def test_failed_tc_collects_capture_and_calls_review_hook(self):
        flow=self.prepare()
        self.host.dump_ui_xml.return_value='<hierarchy><node text="홈"/></hierarchy>'
        run=self.start(flow)
        self.assertEqual(run['counts']['failed'],1)
        case=run['results'][0]
        self.assertTrue(any(e.get('label')=='실패 시점' for e in case['evidence']))
        self.host.record_tc_review_result.assert_called_once()
        self.assertEqual(self.host.record_tc_review_result.call_args.args[1]['status'],'failed')

    def test_standard_format_roundtrip_preserves_manual_result_and_other_rows(self):
        rows=[tc_format.step_row({**step(),'case_no':'DTC-001','depth1':'설정','qa_result':'PASS','automation_result':'FAIL','note':'keep manual note'}),
              tc_format.step_row({**step(),'case_no':'DTC-002','depth1':'설정','qa_result':'FAIL'})]
        book=tc_format.workbook(rows)
        book.create_sheet('Sources').append(['preserve source evidence'])
        raw=tc_format.save(book)
        steps=m.parse_excel(FileStorage(stream=io.BytesIO(raw),filename='standard.xlsx'))
        template=result_workbook.prepare_template(raw,'standard.xlsx',steps)
        self.assertTrue(all(t.automated for t in template.rows))
        self.assertEqual(steps[0]['qa_result'],'PASS')
        self.assertEqual(steps[0]['automation_result'],'FAIL')
        self.assertEqual(steps[0]['depths'][0],'설정')
        self.host._test_workbook=template
        self.host._test_steps=steps
        self.host._test_metadata['file_hash']=template.file_hash
        response=self.call('prepare',{'upload_id':'v1','file_hash':template.file_hash,'indices':[2],'ai_mode':False,'consent':False})
        self.assertEqual(response.status_code,202)
        flow=self.service['current']
        flow.worker.join(3)
        self.assertEqual(self.start(flow)['counts']['passed'],1)
        exported=self.call(flow.id+'/export')
        self.assertEqual(exported.status_code,200)
        result=load_workbook(io.BytesIO(exported.data))
        self.assertEqual(result['TestCase']['L2'].value,'PASS')
        self.assertEqual(result['TestCase']['L3'].value,'FAIL')
        self.assertEqual(result['TestCase']['M2'].value,'FAIL')
        self.assertEqual(result['TestCase']['M3'].value,'PASS')
        self.assertEqual(result['TestCase']['N2'].value,'keep manual note')
        self.assertEqual(result['TestCase']['K3'].value,step()['expected'])
        details=result[result.sheetnames[-1]]
        self.assertEqual(details['A2'].value,'DTC-002')
        self.assertEqual(details['D2'].value,'PASS')
        self.assertTrue(details['F2'].value)
        with zipfile.ZipFile(io.BytesIO(raw)) as before,zipfile.ZipFile(io.BytesIO(exported.data)) as after:
            self.assertEqual(before.read('xl/worksheets/sheet2.xml'),after.read('xl/worksheets/sheet2.xml'))
        result.close()

    def test_generated_excel_headers_parse_and_execute_without_result_column(self):
        book=Workbook()
        book.active.title='TestCase'
        book.active.append(['통합ID','제목','테스트목적','사전조건','테스트절차','기대결과','우선순위','Type','비고'])
        s=step()
        book.active.append([s['case_no'],s['title'],'표시 확인',s['precondition'],s['test_item'],s['expected'],'P1','UI-UX',''])
        buffer=io.BytesIO()
        book.save(buffer)
        book.close()
        self.host._test_steps=m.parse_excel(FileStorage(stream=buffer,filename='reviewed-testcases.xlsx'))
        self.assertEqual(self.host._test_steps[0]['source'],'narrative_qa')
        self.assertEqual(self.host._test_steps[0]['qa_expected'],s['expected'])
        self.assertEqual(self.start(self.prepare())['counts']['passed'],1)

    def test_ai_failure_uses_deterministic_checks_without_second_ai_call(self):
        self.host.request_test_ai_json.side_effect=RuntimeError('private log')
        flow=self.prepare(True)
        self.assertTrue(flow.ai_fallback)
        result=self.start(flow)
        self.assertEqual(result['counts']['passed'],1)
        self.assertTrue(result['ai_fallback'])
        self.host.verify_visual_step.assert_not_called()
        self.assertEqual(self.host.request_test_ai_json.call_count,1)

    def test_precondition_does_not_require_final_screen_labels(self):
        self.host._test_steps[0]['required_labels']=['설정','정보']
        self.host.request_test_ai_json.return_value=({'plans':[ai_plan(self.host._test_steps[0])]}, {})
        self.host.verify_visual_step.side_effect=self.visual
        result=self.start(self.prepare(True))
        self.assertEqual(result['counts']['passed'],1)
        calls=self.host.verify_visual_step.call_args_list
        self.assertEqual(calls[0].args[0]['required_labels'],[])
        self.assertEqual(calls[1].args[0]['required_labels'],[])

    def test_os_precondition_mismatch_never_sends_input(self):
        flow=w.Workflow(self.host,[{**step(),'precondition':'Android 15 단말에서 홈 표시'}],self.metadata,[1],True)
        flow.plans=[{'index':1,'actions':[action()],'blocked_reason':''}]
        flow.execute()
        self.assertEqual(flow.snapshot()['counts']['blocked'],1)
        self.host.verify_visual_step.assert_not_called()
        self.host.adb.assert_not_called()

    def test_ready_or_finished_stop_does_not_cancel_other_execution(self):
        flow=self.prepare()
        flow.stop()
        self.assertFalse(self.host._test_stop_event.is_set())
        flow.state='finished'
        flow.stop()
        self.assertFalse(self.host._test_stop_event.is_set())

    def test_candidate_refresh_rejects_changed_bounds(self):
        flow=w.Workflow(self.host,[step()],self.metadata,[1],True)
        old=self.host.ai_test_screen_candidates.return_value
        moved=copy.deepcopy(old)
        moved[0][0]['bounds']='[30,30][50,50]'
        self.host.ai_test_screen_candidates.side_effect=[old,moved]
        with self.assertRaises(m.BlockedStep):
            flow.action(action('tap','설정'),step())
        self.host.adb.assert_not_called()

    def test_ai_procedure_navigation_and_full_expected(self):
        self.host._test_steps[0]['test_item']='1. 설정 버튼을 클릭한다.'
        self.host.request_test_ai_json.return_value=({'plans':[ai_plan(self.host._test_steps[0],[action('tap','설정')])]}, {})
        self.host.verify_visual_step.side_effect=self.visual
        flow=self.prepare(True)
        result=self.start(flow)
        self.assertEqual(result['counts']['passed'],1)
        self.host.adb.assert_called_once_with('shell','input','tap','10','10',timeout_s=8)
        self.assertEqual(self.host.verify_visual_step.call_args.args[0]['expected'],step()['expected'])
        self.assertEqual(len(result['results'][0]['stages']),4)
        self.assertEqual(result['results'][0]['input_count'],1)

    def test_partial_expected_is_not_pass_and_next_tc_continues(self):
        original=step()
        flow=w.Workflow(self.host,[original,{**original,'case_no':'DTC-002','row':3}],self.metadata,[1,2],True)
        flow.plans=[{'index':i,'actions':[action()],'blocked_reason':''} for i in (1,2)]
        def verify(probe):
            if probe['expected']==step()['expected'] and probe['case_no']=='DTC-001':
                raise AssertionError('필수 기대결과 불충족')
            return self.visual(probe)
        self.host.verify_visual_step.side_effect=verify
        flow.state='running'
        flow.execute()
        self.assertEqual([c['status'] for c in flow.snapshot()['results']],['failed','passed'])

    def test_precondition_failure_never_taps(self):
        flow=self.prepare()
        self.host.dump_ui_xml.return_value='<hierarchy><node package="fixture" text="설정"/></hierarchy>'
        self.assertEqual(self.start(flow)['counts']['blocked'],1)
        self.host.adb.assert_not_called()

    def test_ai_failure_and_missing_data_remain_blocked(self):
        self.host._test_steps[0]['test_item']='1. 강의 화면을 연다.'
        self.host.request_test_ai_json.side_effect=RuntimeError('secret provider log')
        flow=self.prepare(True)
        self.assertIn('AI 계획 실패',flow.plans[0]['blocked_reason'])
        self.assertNotIn('secret',str(flow.public()))
        response=self.call(flow.id+'/start',{'plan_digest':flow.digest,'approved':True})
        self.assertEqual(response.status_code,409)
        self.assertIn('실행 가능한',response.json['error'])
        self.host.begin_redmine_review_run.assert_not_called()
        self.host.adb.assert_not_called()

    def test_risky_environment_reference_and_missing_expected(self):
        for text in ['각 기기에서 확인한다','홈페이지 원문과 비교한다','계정을 삭제한다','앱 재실행을 관찰한다']:
            self.assertTrue(w.limitation({**step(),'test_item':text}))
        self.assertIn('기대결과',w.limitation({**step(),'expected':'','qa_expected':''}))

    def test_omission_reorder_invented_target_and_inspect_shortcut(self):
        s={**step(),'test_item':'1. 설정을 연다.\n2. 뒤로가기를 수행한다.'}
        for actions in [[action('tap','설정')],[action('back',procedure=2),action('tap','설정')],
                        [action(),action('back',procedure=2)],
                        [action('tap','저장'),action('back',procedure=2)],
                        [action('tap','강의'),action('back',procedure=2)]]:
            with self.assertRaises(ValueError):
                w.validate_plan({'index':1,'actions':actions,'blocked_reason':''},s)

    def test_invalid_candidate_and_changed_screen_do_not_send_input(self):
        flow=w.Workflow(self.host,[step()],self.metadata,[1],True)
        for candidates in [[],[{'candidate_id':'other','label':'다른 대상','action':'tap','bounds':'[1,1][20,20]','source':'ui-analysis'}]]:
            self.host.ai_test_screen_candidates.return_value=(candidates,[])
            with self.assertRaises(m.BlockedStep):
                flow.action(action('tap','설정'),step())
        self.host.adb.assert_not_called()

    def test_stop_stale_upload_and_unapproved_plan(self):
        flow=self.prepare()
        self.assertEqual(self.call(flow.id+'/start',{'plan_digest':flow.digest,'approved':False}).status_code,409)
        self.host._test_metadata={**self.metadata,'upload_id':'new'}
        self.assertEqual(self.call(flow.id+'/start',{'plan_digest':flow.digest,'approved':True}).status_code,409)
        self.host._test_metadata=self.metadata
        self.call(flow.id+'/stop',{})
        self.assertEqual(flow.state,'stopped')
        self.assertEqual(self.call(flow.id+'/start',{'plan_digest':flow.digest,'approved':True}).status_code,409)
        self.host.adb.assert_not_called()

    def test_api_consent_selection_and_concurrency(self):
        payload={'upload_id':'v1','file_hash':'hash1','indices':[1],'ai_mode':True,'consent':False}
        self.assertEqual(self.call('prepare',payload).status_code,409)
        self.assertEqual(self.call('prepare',{**payload,'consent':True,'indices':[True]}).status_code,409)
        self.host._random_running=True
        self.assertEqual(self.call('prepare',{**payload,'consent':True}).status_code,409)
        response=self.client.post('/api/tc-workflow/prepare',json=payload,base_url='http://localhost:5000')
        self.assertEqual(response.status_code,403)

    def test_stop_during_navigation_never_passes(self):
        self.host._test_steps[0]['test_item']='1. 설정 버튼을 클릭한다.'
        self.host.request_test_ai_json.return_value=({'plans':[ai_plan(self.host._test_steps[0],[action('tap','설정')])]}, {})
        self.host.verify_visual_step.side_effect=self.visual
        flow=self.prepare(True)
        def stop(*args,**kwargs):
            flow.stop()
            return b'',b'',0
        self.host.adb.side_effect=stop
        self.assertEqual(self.start(flow)['counts']['stopped'],1)

    def modern(self,s,actions,checks):
        plan={'index':1,'actions':actions,'checks':checks,'blocked_reason':''}
        return w.validate_plan(plan,s)

    def test_source_bound_document_and_return_at_different_times(self):
        s={**step(),'precondition':'설정 화면','test_item':'1. 이용약관을 연다.\n2. 내용을 확인한다.',
           'expected':'본문이 표시된다.\n상단 왼쪽 화살표를 누르면 설정으로 돌아온다.'}
        s['qa_expected']=s['expected']
        p=self.modern(s,[{**action('tap','이용약관'),'basis':'procedure','source':'1. 이용약관을 연다.'},
                        {**action(procedure=2),'basis':'procedure','source':'2. 내용을 확인한다.'},
                        {**action('tap','상단 왼쪽 화살표',0),'basis':'expected',
                         'source':'상단 왼쪽 화살표를 누르면 설정으로 돌아온다.'}],
                      [{'expectation':1,'after_action':2},{'expectation':2,'after_action':3}])
        self.host._test_steps=[s]
        screen=['lecture']
        labels={'lecture':('설정',1),'settings':('이용약관',21),'document':('상단 왼쪽 화살표',41)}
        def candidates():
            label,x=labels[screen[0]]
            return ([{'candidate_id':'current','label':label,'bounds':f'[{x},1][{x+19},20]',
                      'action':'tap','source':'ui-analysis'}],[])
        self.host.ai_test_screen_candidates.side_effect=candidates
        def input_(*args,**kwargs):
            screen[0]={'lecture':'settings','settings':'document','document':'settings'}[screen[0]]
            return b'',b'',0
        self.host.adb.side_effect=input_
        def verify(probe):
            expected=probe['expected']
            wanted='document' if expected=='본문이 표시된다.' else 'settings'
            if screen[0]!=wanted:
                raise AssertionError('현재 화면 상태 불일치')
            return self.visual(probe)
        self.host.verify_visual_step.side_effect=verify
        self.host.request_test_ai_json.side_effect=[({'plans':[p]},{}),
            ({'candidate_id':'current','action':'tap','direction':'','reason':'현재 설정 메뉴 선택'}, {})]
        flow=self.prepare(True)
        run=self.start(flow)
        self.assertEqual(run['counts']['passed'],1,run['results'][0]['reason'])
        self.assertEqual(self.host.adb.call_count,3)
        self.assertEqual([x['expectation'] for x in run['results'][0]['stages'] if 'expectation' in x],[1,2])
        self.assertEqual(sum(x.get('supplemental',False) for x in run['results'][0]['stages']),2)
        self.assertNotIn(s['expected'],[c.args[0]['expected'] for c in self.host.verify_visual_step.call_args_list])
        exported=load_workbook(io.BytesIO(self.call(flow.id+'/export').data))
        self.assertEqual(exported.active['M2'].value,'PASS')
        exported.close()

    def test_arbitrary_existing_lecture_resolved_at_runtime(self):
        s={**step(),'precondition':'','test_item':'1. 내 학습에서 임의의 강의 재생 화면을 연다.',
           'expected':'강의 재생 화면이 표시된다.','qa_expected':'강의 재생 화면이 표시된다.'}
        p=self.modern(s,[{**action('navigate','임의의 강의 재생 화면'),
                         'basis':'procedure','source':s['test_item']}],[{'expectation':1,'after_action':1}])
        self.host._test_steps=[s]
        c={'candidate_id':'lecture','label':'기존 강의 A','action':'tap','bounds':'[1,1][20,20]','source':'ui-analysis'}
        self.host.ai_test_screen_candidates.return_value=([c],[])
        self.host.request_test_ai_json.side_effect=[({'plans':[p]},{}),
            ({'candidate_id':'lecture','action':'tap','direction':'','reason':'임의 강의로 현재 기존 강의 선택'}, {})]
        def verify(probe):
            if not self.host.adb.call_count:
                raise AssertionError('강의 화면 진입 전')
            return self.visual(probe)
        self.host.verify_visual_step.side_effect=verify
        run=self.start(self.prepare(True))
        self.assertEqual(run['counts']['passed'],1,run['results'][0]['reason'])
        self.host.adb.assert_called_once()

    def test_split_navigation_stops_at_target_before_later_click(self):
        s={**step(),'precondition':'','test_item':'1. 설정에서 개인정보 처리방침을 클릭한다.',
           'expected':'본문이 표시된다.','qa_expected':'본문이 표시된다.'}
        p=ai_plan(s,[action('navigate','설정'),action('tap','개인정보 처리방침')])
        self.host._test_steps=[s]
        self.host.request_test_ai_json.return_value=({'plans':[p]}, {})
        self.host.ai_test_screen_candidates.return_value=([{'candidate_id':'privacy','label':'개인정보 처리방침',
            'action':'tap','bounds':'[1,1][20,20]','source':'ui-analysis'}],[])
        self.host.verify_visual_step.side_effect=self.visual
        flow=self.prepare(True)
        with patch.object(flow,'navigate',return_value=('현재 목적지 확인',[])) as navigate:
            run=self.start(flow)
        self.assertEqual(run['counts']['passed'],1)
        goals=[call.args[1] for call in navigate.call_args_list]
        self.assertEqual(goals[0],'')
        self.assertTrue(goals[1].startswith('설정 화면이 열려 있다.'))
        self.assertIn('진입 버튼이나 링크만',goals[1])
        self.assertNotIn('개인정보',goals[1])
        self.host.adb.assert_called_once()

    def test_reject_invented_block_and_incomplete_checkpoints(self):
        s=step()
        p={'index':1,'actions':[{**action(),'basis':'procedure','source':s['test_item']}],
           'checks':[{'expectation':1,'after_action':1}],'blocked_reason':''}
        for bad in [{**p,'blocked_reason':'이전 화면과 연동 데이터 없음'}, {**p,'checks':[]},
                    {**p,'checks':[{'expectation':1,'after_action':5}]}]:
            with self.assertRaises(ValueError):
                w.validate_plan(bad,s)
        self.assertEqual(w.limitation({**s,'expected':'홈으로 돌아온다','qa_expected':'홈으로 돌아온다',
                                     'test_purpose':'점수 계산 정확성 데이터 비교'}),'')

    def test_expected_arrow_cannot_be_system_back(self):
        s={**step(),'expected':'화살표를 누르면 설정으로 돌아온다','qa_expected':'화살표를 누르면 설정으로 돌아온다'}
        with self.assertRaises(ValueError):
            self.modern(s,[{**action(),'basis':'procedure','source':s['test_item']},
                           {**action('back',procedure=0),'basis':'expected','source':s['expected']}],
                          [{'expectation':1,'after_action':2}])

    def test_repeated_preparation_action_stops_without_false_pass(self):
        flow=w.Workflow(self.host,[step()],self.metadata,[1],True)
        s={**step(),'_workflow_index':1,'_source_bound':True}
        self.host.verify_visual_step.side_effect=AssertionError('현재 설정 화면 아님')
        self.host.request_test_ai_json.return_value=({'candidate_id':'c001','action':'tap','direction':'','reason':'설정 이동'}, {})
        with self.assertRaisesRegex(m.BlockedStep,'효과 없는'):
            flow.navigate(s,'설정 화면',flow.data['results'][0],True)
        self.assertEqual(self.host.adb.call_count,1)

    def test_navigation_stop_after_ai_response_never_inputs(self):
        flow=w.Workflow(self.host,[step()],self.metadata,[1],True)
        self.host.verify_visual_step.side_effect=AssertionError('현재 설정 화면 아님')
        def stop(*args,**kwargs):
            flow.cancel.set()
            return {'candidate_id':'c001','action':'tap','direction':'','reason':'설정 이동'}, {}
        self.host.request_test_ai_json.side_effect=stop
        with self.assertRaises(w.ai_tc.Stopped):
            flow.navigate({**step(),'_workflow_index':1,'_source_bound':True},'설정 화면',flow.data['results'][0],True)
        self.host.adb.assert_not_called()

    def test_navigation_five_input_budget_and_unavailable(self):
        flow=w.Workflow(self.host,[step()],self.metadata,[1],True)
        s={**step(),'_workflow_index':1,'_source_bound':True}
        self.host.verify_visual_step.side_effect=AssertionError('아직 목표 화면이 아님')
        def candidates():
            n=self.host.adb.call_count
            return ([{'candidate_id':'current','label':'메뉴 '+str(n),'bounds':'[1,1][20,20]',
                      'action':'tap','source':'ui-analysis'}],[])
        self.host.ai_test_screen_candidates.side_effect=candidates
        self.host.request_test_ai_json.return_value=({'candidate_id':'current','action':'tap','direction':'','reason':'메뉴 이동'}, {})
        with self.assertRaisesRegex(m.BlockedStep,'5회 제한'):
            flow.navigate(s,'설정 화면',flow.data['results'][0],True)
        self.assertEqual(self.host.adb.call_count,5)
        self.host.adb.reset_mock()
        self.host.request_test_ai_json.return_value=({'candidate_id':'','action':'unavailable','direction':'','reason':'기존 강의 없음'}, {})
        with self.assertRaisesRegex(m.BlockedStep,'원문:'):
            flow.navigate({**s,'_supplements':0},'임의의 강의 화면',flow.data['results'][0])
        self.host.adb.assert_not_called()

    def test_checkpoint_mismatch_is_fail_not_pass_or_recovery(self):
        s=step()
        p=self.modern(s,[{**action(),'basis':'procedure','source':s['test_item']}],
                      [{'expectation':1,'after_action':1}])
        self.host.request_test_ai_json.return_value=({'plans':[p]}, {})
        def verify(probe):
            if probe['expected']==s['expected']:
                raise AssertionError('기대결과와 실제 화면 불일치')
            return self.visual(probe)
        self.host.verify_visual_step.side_effect=verify
        result=self.start(self.prepare(True))
        self.assertEqual(result['counts']['failed'],1)
        self.host.adb.assert_not_called()
        self.assertEqual(self.host.request_test_ai_json.call_count,1)

    def batch_flow(self,count=2):
        steps=[{**step(),'case_no':f'TC-{i}'} for i in range(1,count+1)]
        return w.Workflow(self.host,steps,self.metadata,list(range(1,count+1)),True)

    def test_malformed_tc_repaired_without_discarding_valid_plan(self):
        flow=self.batch_flow()
        good=[ai_plan(s,index=i) for i,s in enumerate(flow.steps,1)]
        bad=copy.deepcopy(good[1]);bad['actions'][0]['wait_ms']='secret-provider-value'
        self.host.request_test_ai_json.side_effect=[({'plans':[good[0],bad]},{}),({'plans':[good[1]]},{})]
        flow.prepare()
        self.assertEqual(flow.plans,good)
        self.assertEqual(flow.planning[1]['retry_count'],0)
        self.assertEqual(flow.planning[2]['status'],'repaired')
        error=flow.planning[2]['errors'][0]
        self.assertEqual(error['field'],'plan.actions[0].wait_ms')
        self.assertEqual(error['tc_id'],'TC-2')
        self.assertNotIn('secret-provider',str(flow.public()))
        self.assertEqual([c['index'] for c in self.host.request_test_ai_json.call_args.args[1]['cases']],[2])
        self.assertFalse(flow.ai_fallback)
        self.host.adb.assert_not_called()

    def test_failed_repair_only_blocks_affected_tc(self):
        flow=self.batch_flow()
        flow.steps[1]['test_item']='1. 강의 화면을 연다.'
        good=ai_plan(flow.steps[0]);bad=ai_plan(flow.steps[1],[action('navigate','강의 화면')],2)
        bad['checks']=[]
        self.host.request_test_ai_json.side_effect=[({'plans':[good,bad]},{}),({'plans':[bad]},{})]
        flow.prepare()
        self.assertEqual(flow.plans[0],good)
        self.assertTrue(flow.plans[1]['blocked_reason'])
        self.assertEqual(flow.public()['executable_count'],1)
        self.assertEqual(self.host.request_test_ai_json.call_count,2)

    def test_missing_or_duplicate_index_retries_only_affected_tc(self):
        for duplicate in (False,True):
            with self.subTest(duplicate=duplicate):
                flow=self.batch_flow()
                plans=[ai_plan(s,index=i) for i,s in enumerate(flow.steps,1)]
                reply=[plans[0]]+([plans[1],plans[1]] if duplicate else [])
                self.host.request_test_ai_json.reset_mock()
                self.host.request_test_ai_json.side_effect=[({'plans':reply},{}),({'plans':[plans[1]]},{})]
                flow.prepare()
                self.assertEqual(flow.plans,plans)
                self.assertEqual(flow.planning[1]['errors'],[])
                self.assertEqual(flow.planning[2]['errors'][0]['code'],'tc_identity')
                self.assertEqual(self.host.request_test_ai_json.call_count,2)

    def test_broken_batch_retry_consumes_retry_budget(self):
        flow=self.batch_flow()
        plans=[ai_plan(s,index=i) for i,s in enumerate(flow.steps,1)]
        plans[1]['checks']=[]
        self.host.request_test_ai_json.side_effect=[ValueError('private raw JSON'),({'plans':plans},{})]
        flow.prepare()
        self.assertEqual(self.host.request_test_ai_json.call_count,2)
        self.assertEqual(flow.plans[0],plans[0])
        self.assertEqual(flow.planning[2]['retry_count'],1)
        self.assertNotIn('private raw',str(flow.public()))

    def test_transport_errors_never_retry_across_batches(self):
        for code in ('authentication','usage_limit','connection','timeout'):
            with self.subTest(code=code):
                flow=self.batch_flow(11)
                self.host.request_test_ai_json.reset_mock()
                self.host.request_test_ai_json.side_effect=w.AIRequestError(code)
                flow.prepare()
                self.assertEqual(self.host.request_test_ai_json.call_count,1)
                self.assertTrue(all(v['retry_count']==0 for v in flow.planning.values()))
                self.assertTrue(all(v['errors'][0]['code']==code for v in flow.planning.values()))

    def test_stop_during_repair_discards_late_response(self):
        flow=self.batch_flow(1)
        good=ai_plan(flow.steps[0]);bad={**good,'checks':[]}
        def respond(*args,**kwargs):
            if self.host.request_test_ai_json.call_count==1:
                return {'plans':[bad]},{}
            flow.stop()
            return {'plans':[good]},{}
        self.host.request_test_ai_json.side_effect=respond
        flow.prepare()
        self.assertEqual(flow.state,'stopped')
        self.assertEqual(flow.plans,[])
        self.host.adb.assert_not_called()

    def test_planner_can_isolate_schema_errors_but_runtime_validation_is_default(self):
        schema=w.tc_adaptive.SCHEMA
        with patch.object(m,'_test_stop_event',threading.Event()),patch.object(m,'AI_PROVIDER','codex_cli'),patch.object(m,'request_codex_json',return_value=({'plans':[{'index':1}]},{})):
            with self.assertRaises(ValueError):
                m.request_test_ai_json('fixture',{},schema)
            data,_=m.request_test_ai_json('fixture',{},schema,validate_response=False)
            self.assertEqual(data,{'plans':[{'index':1}]})


if __name__=='__main__':
    unittest.main()

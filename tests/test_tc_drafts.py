import io
import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
import urllib.error

from flask import Flask
from openpyxl import load_workbook
from PIL import Image
import app_analysis as a
import figma_import as f
import reference_materials as r
import tc_drafts as d
from local_access import LocalAccessPolicy, HEADER
from test_app_analysis import apk


def result(context):
    eid = next(e['id'] for e in context['evidence'] if e['reference_id'])
    return {'summary': 'Fixture draft; not executed', 'questions': [], 'cases': [
        {'title': 'LC/RC labels', 'objective': 'Verify both labels', 'preconditions': 'Home shown',
         'depth1':'학습 관리','depth2':'실력 분석','depth3':'LC/RC','test_type':'UI-UX',
         'steps': ['Open learning management'], 'expected_results': ['LC and RC are visible'],
         'priority': 'P1', 'category': 'normal', 'source_ids': [eid], 'assumptions': []}]}


def runner(instructions, context, metrics, schema, cancel, stop_type, image):
    metrics['planner_calls'] += 1
    return result(context), {'input_tokens': 12, 'output_tokens': 30}


class DraftTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.store = a.AnalysisStore(Path(self.temp.name))
        self.sha = self.store.register(apk(), 'fixture.apk')[0]['sha256']
        self.refs = r.ReferenceStore(self.store)
        self.ref = self.refs.add(self.sha, io.BytesIO(b'LC and RC are visible.'), 'spec.md', 'requirements')[0]
        self.settings = {'reference_ids': [self.ref['id']], 'scope': 'Learning management', 'count': 2, 'include_images': False}
        self.status = lambda: {'configured': True, 'provider': 'codex_cli', 'model': 'fixture'}
        self.service = d.DraftService(self.store, self.refs, runner, self.status)

    def request(self):
        _, fingerprint, _ = self.service.context(self.sha, self.settings)
        return {**self.settings, 'context_hash': fingerprint, 'consent': True}

    def generate(self):
        batch = self.service.start(self.sha, self.request())
        self.service.worker.join(5)
        self.assertFalse(self.service.worker.is_alive())
        return self.service.get(self.sha, batch['id'])

    def test_generation_snapshot_review_export(self):
        batch = self.generate()
        self.assertEqual(batch['status'], 'completed')
        self.assertEqual(batch['cases'][0]['review_status'], 'draft')
        self.assertEqual(batch['metrics']['planner_calls'], 1)
        self.assertEqual(batch['timeout_seconds'], 600)
        self.assertGreaterEqual(batch['metrics']['elapsed_seconds'], 0)
        self.assertIsNone(batch['error_code'])
        with self.assertRaises(a.AnalysisError):
            self.service.export(self.sha, batch['id'])
        case = batch['cases'][0]
        payload = {k: case[k] for k in d.FIELDS}
        payload.update(review_status='approved', revision=0, title='=HYPERLINK("bad")')
        updated = self.service.review(self.sha, batch['id'], case['id'], payload)
        self.assertEqual(updated['revision'], 1)
        self.assertEqual(updated['original']['cases'][0]['title'], 'LC/RC labels')
        book = load_workbook(io.BytesIO(self.service.export(self.sha, batch['id'])), data_only=False)
        self.assertEqual(book['TestCase']['B2'].data_type, 's')
        self.assertEqual(book['TestCase']['B2'].value, '=HYPERLINK("bad")')
        self.assertEqual(book.sheetnames, ['TestCase', 'Sources'])
        self.assertEqual([c.value for c in book['TestCase'][1]],list(d.tc_format.HEADERS))
        self.assertEqual([book['TestCase'].cell(2,i).value for i in range(4,8)],['학습 관리','실력 분석','LC/RC','UI-UX'])
        self.assertIsNone(book['TestCase']['L2'].value)
        self.assertIsNone(book['TestCase']['M2'].value)
        book.close()
        with self.assertRaises(a.AnalysisError) as caught:
            self.service.review(self.sha, batch['id'], case['id'], payload)
        self.assertEqual(caught.exception.status, 409)
        self.refs.delete(self.sha, self.ref['id'])
        self.assertTrue(self.service.get(self.sha, batch['id'])['context']['evidence'])

    def test_consent_hash_and_missing_provider(self):
        payload = self.request()
        for updates in ({'consent': False}, {'context_hash': 'stale'}):
            with self.assertRaises(a.AnalysisError):
                self.service.start(self.sha, {**payload, **updates})
        self.service.status = lambda: {'configured': False, 'provider': 'codex_cli', 'model': 'fixture'}
        with self.assertRaises(a.AnalysisError) as caught:
            self.service.start(self.sha, payload)
        self.assertEqual(caught.exception.status, 503)
        self.assertEqual(self.service.list(self.sha), [])

    def test_standard_schema_required_and_legacy_draft_preserved(self):
        batch=self.generate()
        case=batch['cases'][0]
        with self.store.connect() as db:
            old_context=dict(batch['context'])
            old_context.pop('tc_format')
            old_case={k:v for k,v in case.items() if k not in d.FORMAT_FIELDS}
            old_case['review_status']='approved'
            db.execute('UPDATE tc_drafts SET context_json=?,cases_json=? WHERE id=?',
                       (json.dumps(old_context),json.dumps([old_case]),batch['id']))
        read=self.service.get(self.sha,batch['id'])
        self.assertEqual(read['cases'][0]['depth1'],'')
        self.assertEqual(read['cases'][0]['test_type'],'')
        self.assertEqual(read['cases'][0]['review_status'],'approved')
        book=load_workbook(io.BytesIO(self.service.export(self.sha,batch['id'])))
        self.assertIsNone(book['TestCase']['D2'].value)
        book.close()
        with self.assertRaises(a.AnalysisError):
            d.validate_case({k:case[k] for k in d.LEGACY_FIELDS},batch['context'])
        with self.assertRaises(a.AnalysisError):
            d.validate_case({**{k:case[k] for k in d.FIELDS},'test_type':''},batch['context'])
        payload={**{k:case[k] for k in d.FIELDS},'review_status':'approved','revision':0}
        self.service.review(self.sha,batch['id'],case['id'],payload)
        legacy={k:v for k,v in payload.items() if k not in d.FORMAT_FIELDS}
        legacy['revision']=1
        saved=self.service.review(self.sha,batch['id'],case['id'],legacy)
        self.assertEqual(saved['cases'][0]['depth1'],'학습 관리')
        self.assertEqual(saved['cases'][0]['test_type'],'UI-UX')

    def test_option_limits_and_cross_build(self):
        for changes in ({'count': True}, {'count': 21}, {'count': '1'}, {'reference_ids': []}, {'scope': ''},
                        {'reference_ids': [self.ref['id']]*2}, {'include_images': 'yes'}):
            with self.assertRaises(a.AnalysisError):
                self.service.context(self.sha, {**self.settings, **changes})
        other = self.store.register(apk({'AndroidManifest.xml': b'other'}), 'other.apk')[0]['sha256']
        with self.assertRaises(a.AnalysisError):
            self.service.context(other, self.settings)

    def test_redaction_and_text_truncation(self):
        with patch.dict(d.os.environ, {'OPENAI_API_KEY': 'private-super-secret'}):
            value = d.redact('private-super-secret user@example.com password=hidden')
        self.assertNotIn('private-super-secret', value)
        self.assertNotIn('user@example.com', value)
        self.assertNotIn('hidden', value)
        ref = self.refs.add(self.sha, io.BytesIO(('A'*12000).encode()), 'long.md', 'requirements')[0]
        context, _, _ = self.service.context(self.sha, {**self.settings, 'reference_ids': [ref['id']]})
        self.assertTrue(any('truncated' in w for w in context['warnings']))

    def test_image_requires_explicit_inclusion(self):
        stream = io.BytesIO()
        Image.new('RGB', (60, 40), 'white').save(stream, 'PNG')
        ref = self.refs.add(self.sha, io.BytesIO(stream.getvalue()), 'design.png', 'design')[0]
        settings = {**self.settings, 'reference_ids': [ref['id']]}
        with self.assertRaises(a.AnalysisError):
            self.service.context(self.sha, settings)
        context, _, images = self.service.context(self.sha, {**settings, 'include_images': True})
        self.assertEqual(len(context['images']), 1)
        self.assertTrue(d.image_sheet(images).startswith(b'\x89PNG'))

    def test_invalid_ai_output_fails_without_fake_cases(self):
        def bad(*args):
            output = result(args[1])
            output['cases'][0]['source_ids'] = ['invented']
            return output, {}
        self.service.runner = bad
        batch = self.generate()
        self.assertEqual(batch['status'], 'failed')
        self.assertIsNone(batch['cases'])
        self.assertEqual(batch['error_code'], 'invalid_response')

    def test_ai_exception_is_sanitized(self):
        def bad(*args):
            raise RuntimeError('secret authentication details')
        self.service.runner = bad
        batch = self.generate()
        self.assertEqual(batch['status'], 'failed')
        self.assertNotIn('secret', json.dumps(batch))
        self.assertEqual(batch['error_code'], 'unknown')

    def test_safe_error_codes_persist_and_old_failure_stays_unknown(self):
        from ai_request_errors import AIRequestError, ERRORS
        for code in ERRORS:
            def fail(*args):
                raise AIRequestError(code)
            self.service.runner = fail
            batch = self.generate()
            self.assertEqual(batch['error_code'], code)
            self.assertEqual(batch['metrics']['error_code'], code)
        with self.store.connect() as db:
            db.execute('UPDATE tc_drafts SET metrics_json=? WHERE id=?', ('{}', batch['id']))
        old = self.service.get(self.sha, batch['id'])
        self.assertIsNone(old['timeout_seconds'])
        self.assertIsNone(old['error_code'])

    def test_stop_discards_late_response_and_busy(self):
        entered = threading.Event()
        def slow(instructions, context, metrics, schema, cancel, stop_type, image):
            entered.set()
            cancel.wait(3)
            return result(context), {}
        self.service.runner = slow
        payload = self.request()
        batch = self.service.start(self.sha, payload)
        self.assertTrue(entered.wait(2))
        with self.assertRaises(a.AnalysisError) as caught:
            self.service.start(self.sha, payload)
        self.assertEqual(caught.exception.status, 409)
        self.service.stop(self.sha, batch['id'])
        self.service.worker.join(5)
        final = self.service.get(self.sha, batch['id'])
        self.assertEqual(final['status'], 'stopped')
        self.assertIsNone(final['cases'])
        self.assertIsNone(final['error_code'])

    def test_restart_marks_pending_interrupted(self):
        with self.store.connect() as db:
            db.execute('INSERT INTO tc_drafts(id,build_sha256,status,created_at,context_json) VALUES (?,?,?,?,?)',
                       ('a'*32, self.sha, 'running', a.now(), '{}'))
        restarted = a.AnalysisStore(self.store.root)
        service = d.DraftService(restarted, self.refs)
        self.assertEqual(service.get(self.sha, 'a'*32)['status'], 'interrupted')

    def test_duplicate_cases_collapsed_and_output_bounds(self):
        context = self.service.context(self.sha, self.settings)[0]
        value = result(context)
        value['cases'] *= 2
        self.assertEqual(len(d.validate_result(value, context)), 1)
        distinct=result(context)
        distinct['cases'].append({**distinct['cases'][0],'depth1':'다른 메뉴'})
        self.assertEqual(len(d.validate_result(distinct,context)),2)
        value['cases'] *= 2
        with self.assertRaises(a.AnalysisError):
            d.validate_result(value, context)

    def test_codex_mode_never_uses_api(self):
        import mirror_server as m
        with patch.object(m, 'AI_PROVIDER', 'codex_cli'), patch.object(m, 'request_codex_json', return_value=({}, {})) as cli, patch.object(m, 'request_openai_json') as api:
            m.request_draft_ai_json('instruction', {}, {}, {}, threading.Event(), d.DraftStopped)
        cli.assert_called_once()
        api.assert_not_called()


class FigmaTests(unittest.TestCase):
    @staticmethod
    def document():
        return {'name': 'Learning', 'version': '123', 'nodes': {'1:2': {'document': {
            'id': '1:2', 'type': 'FRAME', 'name': 'Learning management', 'children': [
                {'id': '1:3', 'type': 'TEXT', 'name': 'LC', 'characters': 'LC results',
                 'absoluteBoundingBox': {'x': 0, 'y': 0, 'width': 50, 'height': 20}}]}}}}

    def test_frame_links_and_rejected_hosts(self):
        key, node, url = f.frame_url('https://www.figma.com/design/abc/Name?node-id=1-2')
        self.assertEqual((key, node), ('abc', '1:2'))
        for value in ['https://www.figma.com/design/abc/Name', 'http://figma.com/design/abc/Name?node-id=1-2',
                      'https://evil.example/design/abc/Name?node-id=1-2', 'https://user@figma.com/design/abc/Name?node-id=1-2',
                      'https://figma.com/design/abc/Name?node-id=1-2&node-id=3-4']:
            with self.assertRaises(a.AnalysisError):
                f.frame_url(value)

    def test_node_normalization_not_runtime_evidence(self):
        value = f.normalize(self.document(), '1:2')
        self.assertEqual(value['nodes'][1]['text'], 'LC results')
        self.assertEqual(value['version'], '123')
        for data in [{}, {'nodes': []}, {'nodes': {'1:2': None}}, {**self.document(), 'version': ''}]:
            with self.assertRaises(a.AnalysisError):
                f.normalize(data, '1:2')

    def test_network_credential_host_boundary(self):
        for url, token in [('http://api.figma.com/v1/x', 'secret'), ('https://evil.example/x', 'secret'),
                           ('https://api.figma.com/v1/x', None), ('https://s3-alpha.figma.com/img', 'secret')]:
            with self.assertRaises(a.AnalysisError):
                f.fetch(url, token)
        class Reply:
            def __enter__(self): return self
            def __exit__(self, *args): pass
            def read(self, size): return b'{}'
        with patch.object(f.urllib.request, 'build_opener') as build:
            build.return_value.open.return_value = Reply()
            f.fetch('https://api.figma.com/v1/files/key', 'secret')
            req = build.return_value.open.call_args.args[0]
            self.assertIn('secret', req.headers.values())
            f.fetch('https://s3-alpha.figma.com/img')
            req = build.return_value.open.call_args.args[0]
            self.assertNotIn('secret', req.headers.values())

    def test_auth_permission_rate_limit_and_size(self):
        for code in (401, 403, 404, 429, 302):
            with patch.object(f.urllib.request, 'build_opener') as build:
                build.return_value.open.side_effect = urllib.error.HTTPError('url', code, 'private details', {}, None)
                with self.assertRaises(a.AnalysisError) as caught:
                    f.fetch('https://api.figma.com/v1/x', 'secret')
                self.assertEqual(caught.exception.status, code if code != 302 else 502)
                self.assertNotIn('private details', str(caught.exception))

    def test_import_missing_token_and_partial_image(self):
        with tempfile.TemporaryDirectory() as folder:
            store = a.AnalysisStore(folder)
            sha = store.register(apk(), 'a.apk')[0]['sha256']
            refs = r.ReferenceStore(store)
            url = 'https://www.figma.com/design/abc/Name?node-id=1-2'
            with patch.dict(f.os.environ, {'FIGMA_ACCESS_TOKEN': ''}):
                with self.assertRaises(a.AnalysisError) as caught:
                    f.import_frame(refs, sha, url)
                self.assertEqual(caught.exception.status, 503)
            with patch.dict(f.os.environ, {'FIGMA_ACCESS_TOKEN': 'secret'}), patch.object(f, 'api_json', side_effect=[self.document(), {'images': {}}]):
                item, duplicate = f.import_frame(refs, sha, url)
            self.assertFalse(duplicate)
            self.assertEqual(item['preview']['status'], 'partial')
            self.assertNotIn('secret', json.dumps(item))


class DraftApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.app = Flask('drafts-test')
        self.policy = LocalAccessPolicy('127.0.0.1', 5000, ['localhost', '127.0.0.1'])
        self.policy.install(self.app)
        a.register(self.app, self.policy, self.temp.name, runner, lambda: {'configured': True, 'provider': 'codex_cli', 'model': 'fixture'})
        self.service = self.app.extensions['tc_drafts']
        self.sha = self.service.store.register(apk(), 'a.apk')[0]['sha256']
        ref = self.service.refs.add(self.sha, io.BytesIO(b'LC and RC visible'), 'a.md', 'requirements')[0]
        self.settings = {'reference_ids': [ref['id']], 'scope': 'Labels', 'count': 1, 'include_images': False}
        self.base = '/api/app-analysis/builds/' + self.sha
        self.client = self.app.test_client()
        self.headers = {HEADER: self.policy.token, 'Origin': 'http://localhost:5000'}

    def call(self, path, **kwargs):
        kwargs.setdefault('headers', self.headers)
        return self.client.open(self.base + path, base_url='http://localhost:5000', **kwargs)

    def test_context_generate_list_detail_and_unknown_reference(self):
        preview = self.call('/draft-context', method='POST', json=self.settings)
        self.assertEqual(preview.status_code, 200)
        res = self.call('/drafts', method='POST', json={**self.settings, 'context_hash': preview.json['context_hash'], 'consent': True})
        self.assertEqual(res.status_code, 202)
        self.service.worker.join(5)
        job = res.json['draft']['id']
        self.assertEqual(self.call('/drafts/' + job).json['draft']['status'], 'completed')
        self.assertEqual(len(self.call('/drafts').json['drafts']), 1)
        self.assertEqual(self.call('/drafts/'+job+'/export').status_code, 409)
        self.assertEqual(self.call('/drafts/unknown').status_code, 404)

    def test_origin_token_and_extra_payload_rejected(self):
        for headers in [{}, {**self.headers, 'Origin': 'http://evil.example'}, {**self.headers, 'X-Forwarded-For': '127.0.0.1'}]:
            self.assertEqual(self.call('/draft-context', method='POST', headers=headers, json=self.settings).status_code, 403)
        self.assertEqual(self.call('/draft-context', method='POST', json={**self.settings, 'api_key': 'bad'}).status_code, 400)
        self.assertEqual(self.call('/figma', method='POST', json={'url': 'https://figma.com', 'token': 'bad'}).status_code, 400)

    def test_standard_template_guard_and_blank_results(self):
        path='/api/app-analysis/testcase-template'
        self.assertEqual(self.client.get(path,base_url='http://localhost:5000').status_code,403)
        response=self.client.get(path,base_url='http://localhost:5000',headers=self.headers)
        self.assertEqual(response.status_code,200)
        book=load_workbook(io.BytesIO(response.data))
        self.assertEqual(tuple(c.value for c in book.active[1]),d.tc_format.HEADERS)
        self.assertEqual(book.active.max_row,1)
        book.close()
        self.assertEqual(self.service.list(self.sha),[])


if __name__ == '__main__':
    unittest.main()

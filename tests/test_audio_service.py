import io
import json
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from flask import Flask
import audio_service as service


HTML = '<span class="playlist_text">Step 1. 단어 학습</span><div class="playlist-item" audio_url="seoch.php?Id=1&amp;app=jp"><div class="playlist-subtitle">hello</div></div><span class="playlist_text">Step 3. 회화 학습</span><div class="playlist-item" audio_url="seoch.php?Id=2&amp;app=jp"><div class="playlist-subtitle">different</div></div>'
UNITS = [dict(key=f'L2-W3-D{d}', level=2, week=3, day=d, url=service.BASE+f'jday.php?l=2&w=3&d={d}') for d in (1, 3)]


class AudioServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.app = Flask(__name__)
        self.state = service.register(self.app, lambda: SimpleNamespace(capture_screen_png=lambda: b'png'), self.root)
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp.cleanup()

    def wait(self):
        for _ in range(200):
            d = self.client.get('/api/audio-verification').json['current']
            if d['status'] != 'running':
                return d
            time.sleep(.02)
        self.fail('worker timed out')

    def prepare(self):
        with patch.object(service, 'discover_catalog', return_value=UNITS), patch.object(service, 'fetch', return_value=HTML.encode()):
            return self.client.post('/api/audio-verification/preview', json={'units': [u['key'] for u in UNITS]}).json

    def test_parse_stages_not_fixed_positions(self):
        rows = service.parse_items(HTML, UNITS[0]['url'])
        self.assertEqual([x['stage'] for x in rows], ['word', 'conversation'])
        self.assertEqual(len(rows), 2)

    def test_catalog_uses_actual_buttons_only(self):
        pages = [b'<div onclick="func.goWeek(2)"></div>', b'''<div onclick="window.location.href='jday.php?l=2&w=3&d=3'"></div><a href="https://evil.test/jday.php?l=2&w=3&d=4">bad</a>''']
        with patch.object(service, 'fetch', side_effect=pages):
            self.assertEqual([u['key'] for u in service.discover_catalog()], ['L2-W3-D3'])

    def test_invalid_and_duplicate_selection(self):
        for keys in [[], ['a', 'a'], [1]]:
            self.assertEqual(self.client.post('/api/audio-verification/preview', json={'units': keys}).status_code, 400)
        with patch.object(service, 'discover_catalog', return_value=UNITS):
            self.assertEqual(self.client.post('/api/audio-verification/preview', json={'units':['L999-W1-D1']}).status_code, 400)

    def test_batch_evidence_export_and_isolation(self):
        p = self.prepare()
        model = SimpleNamespace(transcribe=lambda *a, **k: ([SimpleNamespace(text='hello')], SimpleNamespace(duration=3)))
        with patch.object(service, 'fetch', return_value=b'audio'), patch.object(service, 'load_model', return_value=model):
            r = self.client.post('/api/audio-verification', json={'preview_id':p['preview_id']})
            self.assertEqual(r.status_code, 202)
            self.assertEqual(self.client.post('/api/audio-verification', json={'preview_id':p['preview_id']}).status_code, 409)
            d = self.wait()
        self.assertEqual(d['total'], 4)
        self.assertEqual(d['counts'], {'match':2, 'review':2})
        for u in UNITS:
            folder = self.root/'japanese'/d['id']/u['key']
            self.assertTrue((folder/'01.mp3').exists())
            self.assertTrue((folder/'02.txt').exists())
            self.assertTrue((folder/'source.html').exists())
            self.assertEqual(json.loads((folder/'result.json').read_text(encoding='utf-8'))['key'], u['key'])
        with self.client.get(f"/api/audio-verification/{d['id']}/L2-W3-D3/01.mp3") as response:
            self.assertEqual(response.data, b'audio')
        self.assertEqual(self.client.get(f"/api/audio-verification/{d['id']}/L2-W3-D3/source.html").status_code, 404)
        self.assertEqual(self.client.get(f"/api/audio-verification/{d['id']}/../result.json").status_code, 404)
        from openpyxl import load_workbook
        r = self.client.get(f"/api/audio-verification/{d['id']}/export")
        wb = load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb['음원 검증'].max_row, 5)
        self.assertEqual(wb['단원 요약'].max_row, 3)
        self.assertEqual(wb['음원 검증']['F2'].value, '단어 학습')
        self.assertIn('화면 문구와 일치', wb['음원 검증']['J2'].value)
        self.assertEqual(wb['단원 요약']['B2'].value, '검사 완료')

    def test_model_failure_and_resume(self):
        p = self.prepare()
        with patch.object(service, 'load_model', side_effect=RuntimeError('secret')):
            self.client.post('/api/audio-verification', json={'preview_id':p['preview_id']})
            d = self.wait()
        self.assertEqual(d['status'], 'failed')
        self.assertNotIn('secret', d['error'])
        model = SimpleNamespace(transcribe=lambda *a, **k: ([SimpleNamespace(text='hello')], SimpleNamespace(duration=3)))
        with patch.object(service, 'load_model', return_value=model), patch.object(service, 'fetch', return_value=b'audio'):
            self.assertEqual(self.client.post('/api/audio-verification/resume').status_code, 202)
            resumed = self.wait()
        self.assertNotEqual(d['id'], resumed['id'])
        self.assertEqual(resumed['resumed_from'], d['id'])
        self.assertEqual(json.loads((self.root/'japanese'/d['id']/'result.json').read_text(encoding='utf-8'))['status'], 'failed')

    def test_premium_sessions_and_subject_audio_validation(self):
        pages=[b'<div onclick="func.goWeek(1)"></div>', b'<a href="tjp2day.php?l=1&w=1&d=1&s=1">one</a><a href="tjp2day.php?l=1&w=1&d=1&s=2">two</a>']
        with patch.object(service,'fetch',side_effect=pages):
            units=service.discover_catalog('japanese-premium')
        self.assertEqual([u['key'] for u in units],['L1-W1-D1-S1','L1-W1-D1-S2'])
        with self.assertRaises(ValueError): service.parse_items(HTML,units[0]['url'],'japanese-premium')
        rows=service.parse_items(HTML.replace('app=jp','app=tjp2'),units[0]['url'],'japanese-premium')
        self.assertEqual(len(rows),2)

    def test_premium_storage_history_and_cross_subject_preview(self):
        unit=dict(key='L1-W1-D1-S1',level=1,week=1,day=1,session=1,url=service.BASE+'tjp2day.php?l=1&w=1&d=1&s=1')
        with patch.object(service,'discover_catalog',return_value=[unit]),patch.object(service,'fetch',return_value=HTML.replace('app=jp','app=tjp2').encode()):
            p=self.client.post('/api/audio-verification/preview',json={'subject':'japanese-premium','units':[unit['key']]}).json
        self.assertEqual(self.client.post('/api/audio-verification',json={'subject':'japanese','preview_id':p['preview_id']}).status_code,409)
        model=SimpleNamespace(transcribe=lambda *a,**k:([SimpleNamespace(text='hello')],SimpleNamespace(duration=3)))
        with patch.object(service,'fetch',return_value=b'audio'),patch.object(service,'load_model',return_value=model):
            self.client.post('/api/audio-verification',json={'subject':'japanese-premium','preview_id':p['preview_id']})
            d=self.wait()
        self.assertTrue((self.root/'japanese-premium'/d['id']/unit['key']/'01.mp3').is_file())
        self.assertEqual(d['subject'],'japanese-premium')
        self.assertEqual(self.client.get('/api/audio-verification/history?subject=japanese').json['runs'],[])
        self.assertEqual(len(self.client.get('/api/audio-verification/history?subject=japanese-premium').json['runs']),1)
        with self.client.get(f"/api/audio-verification/{d['id']}/{unit['key']}/01.mp3") as response:
            self.assertEqual(response.data,b'audio')

    def test_legacy_history_is_read_without_rewriting(self):
        rid='b'*32; folder=self.root/rid; folder.mkdir()
        raw=json.dumps({'id':rid,'status':'finished','items':[],'units':[]})
        (folder/'result.json').write_text(raw,encoding='utf-8')
        rows=self.client.get('/api/audio-verification/history?subject=japanese').json['runs']
        self.assertEqual(rows[0]['subject'],'japanese')
        self.assertEqual(self.client.get(f'/api/audio-verification/{rid}/result.json').json['subject'],'japanese')
        self.assertEqual((folder/'result.json').read_text(encoding='utf-8'),raw)

    def test_unknown_subject_rejected(self):
        self.assertEqual(self.client.get('/api/audio-verification/catalog?subject=unknown').status_code,400)
        self.assertEqual(self.client.post('/api/audio-verification/preview',json={'subject':'unknown','units':['L1-W1-D1']}).status_code,400)

    def test_stop_discards_inflight_and_resume(self):
        p = self.prepare()
        def transcribe(*a, **k):
            self.client.post('/api/audio-verification/stop')
            return [SimpleNamespace(text='hello')], SimpleNamespace(duration=3)
        with patch.object(service, 'load_model', return_value=SimpleNamespace(transcribe=transcribe)), patch.object(service, 'fetch', return_value=b'audio'):
            self.client.post('/api/audio-verification', json={'preview_id':p['preview_id']})
            d = self.wait()
        self.assertEqual(d['status'], 'stopped')
        self.assertEqual(d['items'], [])

    def test_network_failure_continues(self):
        p = self.prepare()
        with patch.object(service, 'load_model', return_value=object()), patch.object(service, 'fetch', side_effect=OSError()):
            self.client.post('/api/audio-verification', json={'preview_id':p['preview_id']})
            d = self.wait()
        self.assertEqual(d['counts'], {'unavailable': 4})

    def test_fetch_rejects_other_hosts(self):
        for url in ['https://evil.test/jday.php', 'http://theweek.brain-study.co.kr/jday.php',service.BASE+'admin.php']:
            with self.assertRaises(ValueError): service.fetch(url, 10)

    def test_empty_day_is_not_silently_counted_as_pass(self):
        with patch.object(service, 'discover_catalog', return_value=UNITS), patch.object(service, 'fetch', return_value=b'<html></html>'):
            p = self.client.post('/api/audio-verification/preview', json={'units':[UNITS[0]['key']]}).json
        self.assertEqual(p['total'], 0)
        self.assertEqual(p['units'][0]['status'], 'unavailable')
        self.assertEqual(self.client.post('/api/audio-verification', json={'preview_id':p['preview_id']}).status_code, 400)

    def test_restart_restores_history_without_execution(self):
        rid = 'a' * 32
        (self.root/rid).mkdir()
        (self.root/rid/'result.json').write_text(json.dumps({'id':rid,'status':'running','items':[]}),encoding='utf-8')
        app = Flask('restore')
        with patch.object(service, 'load_model', side_effect=AssertionError('must not run')):
            restored = service.register(app, lambda: None, self.root)
        self.assertEqual(restored['current']['status'], 'stopped')
        self.assertEqual(restored['current']['id'], rid)


if __name__ == '__main__':
    unittest.main()

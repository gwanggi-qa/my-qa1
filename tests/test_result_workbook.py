"""Synthetic workbooks and mocked runs only; never operates a connected tablet."""
import copy
from contextlib import ExitStack
from dataclasses import replace
import hashlib
import io
import unittest
import zipfile
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import PatternFill
from PIL import Image

import mirror_server as m
import result_workbook as rw


def fixture(full=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCase"
    ws.merge_cells("A1:K1")
    ws["A1"] = "Synthetic QA report"
    for col, text in {2: "코드번호", 3: "과목명", 4: "No.", 8: "전제조건",
                      9: "테스트 항목", 10: "기대결과", 11: "검증결과\n(P,F,NT,NA)"}.items():
        ws.cell(4, col, text)
    steps = []
    for row in range(6, 295 if full else 12):
        number = row - 5 if row <= 20 else row + 25
        case = f"TSCL-{number:03d}"
        ws.cell(row, 2, "code-16")
        ws.cell(row, 3, "기초영어(TC)")
        ws.cell(row, 4, case)
        ws.cell(row, 9, "Home 버튼 클릭\nReport 버튼 클릭")
        ws.cell(row, 10, "Report 표시")
        ws.cell(row, 11).fill = PatternFill("solid", fgColor="FFFF00")
        steps.append({"sheet": "TestCase", "row": row, "case_no": case, "action": "tap",
                      "source": "wordp_testcase", "test_item": ws.cell(row, 9).value, "target": "Home"})
    ws["K6"].comment = Comment("Preserve original comment", "Fixture")
    png = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(png, "PNG")
    png.seek(0)
    ws.add_image(ExcelImage(png), "M1")
    old = wb.create_sheet("TC(작업 중)")
    old.append(["No", "검증결과"])
    old.append(["OLD-1", "p"])
    old.append(["OLD-2", "nt"])
    history = wb.create_sheet("History")
    history["P2"] = "Mate L2E"
    history["Q2"] = "주니어영어(TC)"
    history["R2"] = "l2e-*"
    history["P4"], history["Q4"], history["R4"], history["S4"] = "P", "F", "NT", "NA"
    history["P5"] = '=COUNTIFS(TestCase!$D:$D,$R$2,TestCase!$K:$K,P4)'
    history["Q5"] = '=COUNTIFS(TestCase!$D:$D,$R$2,TestCase!$K:$K,Q4)'
    history["O5"] = '=COUNTIF(TestCase!D:D,R2)'
    history["D9"] = "왕초보스피킹R"
    for row in (9, 10, 60):
        history[f"I{row}"] = f'=COUNTIFS(TestCase!$C:$C,$D{row},TestCase!$K:$K,P$4)'
        history[f"J{row}"] = f'=COUNTIFS(TestCase!$C:$C,$D{row},TestCase!$K:$K,Q$4)'
        history[f"H{row}"] = f'=SUM(I{row}:L{row})'
    data = io.BytesIO()
    wb.save(data)
    return data.getvalue(), steps


def read(payload):
    return load_workbook(io.BytesIO(payload), data_only=False)


class WorkbookExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.payload, cls.steps = fixture()
        cls.template = rw.prepare_template(cls.payload, "source.xlsx", cls.steps)

    def export(self, results=(), stopped=False, active=()):
        run = rw.finish_run("a" * 32, self.template, self.steps, results, stopped, active)
        return rw.export_run(run, self.template.file_hash)

    def test_results_are_exact_codes_and_unfinished_never_pass(self):
        results = [{"index": i, "status": status} for i, status in enumerate(
            ["passed", "failed", "blocked", "stopped", "not_applicable"], 1)]
        wb = read(self.export(results))
        self.assertEqual([wb["TestCase"][f"K{r}"].value for r in range(6, 12)], ["Pass", "F", "NT", "NT", "NA", "NT"])
        self.assertEqual(wb["TC(작업 중)"]["B2"].value, "p")
        self.assertEqual(wb["TC(작업 중)"]["B3"].value, "nt")
        wb.close()

    def test_four_success_formulas_only_and_filters_unchanged(self):
        self.assertEqual([(s, c) for s, c, _ in self.template.formulas],
                         [("History", "P5"), ("History", "I9"), ("History", "I10"), ("History", "I60")])
        wb = read(self.export())
        before = read(self.payload)
        for cell in ("P5", "I9", "I10", "I60"):
            value = wb["History"][cell].value
            self.assertIn('"P"', value)
            self.assertIn('"Pass"', value)
            self.assertNotIn('"P*"', value)
            self.assertEqual(value.count("COUNTIFS("), 2)
            self.assertNotIn("P$4", value)
            self.assertNotIn(",P4", value)
        for cell in ("P2", "Q2", "R2", "D9", "D10", "D60", "Q5", "O5", "H9", "H10", "H60", "J9"):
            self.assertEqual(wb["History"][cell].value, before["History"][cell].value)
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertTrue(wb.calculation.forceFullCalc)
        before.close()
        wb.close()

    def test_formula_compatibility_is_idempotent_and_header_independent(self):
        exported = self.export()
        template = rw.prepare_template(exported, "again.xlsx", self.steps)
        self.assertEqual(template.formulas, ())
        wb = read(self.payload)
        wb["History"]["P4"] = "Pass"
        cols = {"TestCase": {"K"}}
        formula = wb["History"]["I9"].value
        result = rw.compatible_formula(formula, "History", wb, cols)
        self.assertEqual(result, '=(COUNTIFS(TestCase!$C:$C,$D9,TestCase!$K:$K,"P")+COUNTIFS(TestCase!$C:$C,$D9,TestCase!$K:$K,"Pass"))')
        self.assertEqual(rw.compatible_formula(result, "History", wb, cols), result)
        wb.close()

    def test_formula_ranges_and_non_success_criteria_are_not_guessed(self):
        wb = read(self.payload)
        cols = {"TestCase": {"K"}}
        unchanged = ['=COUNTIF(TestCase!K:K,"Pending")', '=COUNTIF(TestCase!K:K,"P*")',
                     '=COUNTIF(TestCase!C:C,"P")', '=COUNTIF(TestCase!K:K,"F")',
                     '=SUMPRODUCT((TestCase!K6:K11="P")*1)', '="COUNTIF(TestCase!K:K,""P"")"']
        for formula in unchanged:
            self.assertEqual(rw.compatible_formula(formula, "History", wb, cols), formula)
        formula = '=COUNTIF(\'TestCase\'!$K$6:$K$11,"p")'
        changed = rw.compatible_formula(formula, "History", wb, cols)
        self.assertIn('"P"', changed)
        self.assertIn('"Pass"', changed)
        wb.close()

    def test_mixed_legacy_and_pass_values_count_once_without_pending(self):
        wb = read(self.payload)
        original = '=COUNTIF(TestCase!K6:K11,"P")'
        changed = rw.compatible_formula(original, "History", wb, {"TestCase": {"K"}})
        calls = list(rw._count_calls(changed))
        criteria = [args[1].strip('"') for _, _, _, args in calls]
        self.assertEqual(criteria, ["P", "Pass"])
        values = ["P", "p", "Pass", "PASS", "Pending", "F", "NT", "NA", None]
        count = sum(sum(str(value or "").casefold() == criterion.casefold() for value in values) for criterion in criteria)
        self.assertEqual(count, 4)
        wb.close()

    def test_package_other_members_styles_merges_images_comments_preserved(self):
        exported = self.export([{"index": 1, "status": "passed"}])
        self.assertEqual(hashlib.sha256(self.payload).hexdigest(), self.template.file_hash)
        with zipfile.ZipFile(io.BytesIO(self.payload)) as original, zipfile.ZipFile(io.BytesIO(exported)) as output:
            self.assertEqual(original.namelist(), output.namelist())
            changes = {name for name in original.namelist() if original.read(name) != output.read(name)}
            self.assertEqual(changes, {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet3.xml", "xl/workbook.xml"})
        wb = read(exported)
        self.assertEqual(wb["TestCase"]["K6"].fill.fgColor.rgb, "00FFFF00")
        self.assertEqual(wb["TestCase"]["K6"].comment.text, "Preserve original comment")
        self.assertIn("A1:K1", str(wb["TestCase"].merged_cells))
        self.assertEqual(len(wb["TestCase"]._images), 1)
        wb.close()

    def test_all_289_rows_keep_nonconsecutive_original_ids(self):
        payload, steps = fixture(full=True)
        template = rw.prepare_template(payload, "report.xlsx", steps)
        self.assertEqual(len(template.rows), 289)
        self.assertEqual([(t.row, t.case_no) for t in template.rows[14:17]], [(20, "TSCL-015"), (21, "TSCL-046"), (22, "TSCL-047")])
        run = rw.finish_run("b" * 32, template, steps, [{"index": 16, "status": "passed"}])
        wb = read(rw.export_run(run, template.file_hash))
        self.assertEqual(wb["TestCase"]["K21"].value, "Pass")
        self.assertEqual(wb["TestCase"]["D294"].value, "TSCL-319")
        self.assertEqual(wb["TestCase"]["K294"].value, "NT")
        self.assertEqual(wb["TestCase"]["K20"].value, "NT")
        wb.close()

    def test_hash_row_id_duplicate_and_unknown_status_fail_closed(self):
        run = rw.finish_run("c" * 32, self.template, self.steps, [])
        with self.assertRaises(ValueError):
            rw.export_run(run, "0" * 64)
        with self.assertRaises(ValueError):
            rw.export_run(replace(run, template=replace(self.template, payload=b"tampered")), self.template.file_hash)
        steps = copy.deepcopy(self.steps)
        steps[0]["case_no"] = "WRONG"
        with self.assertRaises(ValueError):
            rw.finish_run("c" * 32, self.template, steps, [])
        for results in ([{"index": 1, "status": "running"}], [{"index": 999, "status": "passed"}],
                        [{"index": 1, "status": "passed"}, {"index": 1, "status": "failed"}]):
            with self.assertRaises(ValueError):
                rw.finish_run("c" * 32, self.template, self.steps, results)
        with self.assertRaises(ValueError):
            rw.prepare_template(self.payload, "source.xlsx", self.steps + self.steps[:1])

    def test_merged_or_formula_result_cells_are_rejected(self):
        for merge in (True, False):
            wb = read(self.payload)
            if merge:
                wb["TestCase"].merge_cells("K6:K7")
            else:
                wb["TestCase"]["K6"] = '=IF(A6="","","P")'
            output = io.BytesIO()
            wb.save(output)
            wb.close()
            with self.assertRaises(ValueError):
                rw.prepare_template(output.getvalue(), "source.xlsx", self.steps)

    def test_missing_result_cell_is_inserted_in_column_order(self):
        with zipfile.ZipFile(io.BytesIO(self.payload)) as original:
            out = io.BytesIO()
            with zipfile.ZipFile(out, "w") as dest:
                for info in original.infolist():
                    data = original.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        xml = data.decode()
                        match = rw._cell_match(xml, "K6")
                        xml = xml[:match.start()] + xml[match.end():]
                        data = xml.encode()
                    dest.writestr(info, data)
        payload = out.getvalue()
        template = rw.prepare_template(payload, "source.xlsx", self.steps)
        run = rw.finish_run("d" * 32, template, self.steps, [{"index": 1, "status": "passed"}])
        wb = read(rw.export_run(run, template.file_hash))
        self.assertEqual(wb["TestCase"]["K6"].value, "Pass")
        wb.close()


class WorkbookRouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.payload, cls.steps = fixture()
        cls.template = rw.prepare_template(cls.payload, "source.xlsx", cls.steps)

    def setUp(self):
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        for name, value in [("_test_workbook", self.template), ("_last_workbook_run", None),
                            ("_test_steps", copy.deepcopy(self.steps)), ("_test_running", False),
                            ("_random_running", False), ("_record_play_running", False),
                            ("_test_metadata", {"upload_id": "upload-one", "file_hash": self.template.file_hash, "filename": "source.xlsx"})]:
            self.stack.enter_context(patch.object(m, name, value))
        m._test_stop_event.clear()
        self.addCleanup(m._test_stop_event.clear)
        # Fail immediately if a test accidentally reaches an external side effect.
        for name in ("adb", "request_test_ai_json", "redmine_api_json"):
            self.stack.enter_context(patch.object(m, name, side_effect=AssertionError("External action forbidden")))
        self.client = m.app.test_client()

    def complete(self):
        snapshot = rw.finish_run("e" * 32, self.template, self.steps, [{"index": 1, "status": "passed"}])
        m._last_workbook_run = snapshot
        return f"/api/test-runs/{snapshot.run_id}/workbook?file_hash={self.template.file_hash}"

    def upload(self, **fields):
        return self.client.post("/api/testcase", data={"file": (io.BytesIO(self.payload), "source.xlsx"), **fields}, content_type="multipart/form-data")

    def test_download_real_uploaded_template_copy_and_local_only(self):
        url = self.complete()
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertIn("attachment", response.headers["Content-Disposition"])
        wb = read(response.data)
        self.assertEqual(wb["TestCase"]["K6"].value, "Pass")
        wb.close()
        self.assertEqual(self.client.get(url, environ_base={"REMOTE_ADDR": "192.168.1.20"}).status_code, 403)

    def test_wrong_run_hash_running_and_reupload_reject_old_download(self):
        url = self.complete()
        self.assertEqual(self.client.get(url.replace("e" * 32, "f" * 32)).status_code, 409)
        self.assertEqual(self.client.get(url.replace(self.template.file_hash, "0" * 64)).status_code, 409)
        with patch.object(m, "_test_running", True):
            self.assertEqual(self.client.get(url).status_code, 409)
            self.assertEqual(self.upload().status_code, 409)
        self.assertEqual(self.upload().status_code, 200)
        self.assertIsNone(m._last_workbook_run)
        self.assertEqual(self.client.get(url).status_code, 409)

    def test_preview_and_failed_upload_keep_existing_snapshot(self):
        url = self.complete()
        snapshot = m._last_workbook_run
        self.assertEqual(self.upload(preview="true").status_code, 200)
        self.assertIs(m._last_workbook_run, snapshot)
        response = self.client.post("/api/testcase", data={"file": (io.BytesIO(b"broken"), "broken.xlsx")})
        self.assertEqual(response.status_code, 400)
        self.assertIs(m._last_workbook_run, snapshot)
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_reupload_during_download_generation_rejects_stale_response(self):
        url = self.complete()
        original = rw.export_run
        def replace_while_exporting(snapshot, file_hash):
            data = original(snapshot, file_hash)
            m._last_workbook_run = None
            return data
        with patch.object(rw, "export_run", side_effect=replace_while_exporting):
            self.assertEqual(self.client.get(url).status_code, 409)

    def test_replacement_during_preflight_cannot_start_other_workbook(self):
        def replace_upload():
            m._test_metadata = {**m._test_metadata, "upload_id": "upload-two"}
            return True
        with patch.object(m, "check_device", side_effect=replace_upload), patch.object(m, "validate_recorder_resolution"), patch.object(m.threading, "Thread") as thread:
            self.assertEqual(self.client.post("/api/run", json={}).status_code, 409)
        thread.assert_not_called()

    def test_workbook_raw_run_requires_review(self):
        with patch.object(m, "check_device", return_value=True), patch.object(m, "validate_recorder_resolution"), patch.object(m.threading, "Thread") as thread:
            response = self.client.post("/api/run", json={})
        self.assertEqual(response.status_code, 409)
        thread.assert_not_called()

    def test_mock_run_stopped_row_and_remaining_rows_are_nt(self):
        count = 0
        def execute(step):
            nonlocal count
            count += 1
            if count == 2:
                m._test_stop_event.set()
                raise m.ai_tc.Stopped("stopped")
            return "Mock verified"
        with ExitStack() as stack:
            for name in ("validate_recorder_resolution", "record_linked_retest", "append_test_run_history", "append_ai_test_history"):
                stack.enter_context(patch.object(m, name))
            stack.enter_context(patch.object(m, "check_device", return_value=True))
            stack.enter_context(patch.object(m, "execute_standard_test_step", side_effect=execute))
            emit = stack.enter_context(patch.object(m.socketio, "emit"))
            m.run_test_steps()
        snapshot = m._last_workbook_run
        self.assertEqual(snapshot.statuses, ("passed", "stopped", "not_run", "not_run", "not_run", "not_run"))
        finished = [c.args[1] for c in emit.call_args_list if c.args[0] == "test_run_finished"][0]
        self.assertTrue(finished["stopped"])
        self.assertEqual(finished["workbook_export"]["run_id"], snapshot.run_id)
        wb = read(rw.export_run(snapshot, self.template.file_hash))
        self.assertEqual([wb["TestCase"][f"K{r}"].value for r in range(6, 12)], ["Pass", "NT", "NT", "NT", "NT", "NT"])
        wb.close()

    def test_environment_failure_before_first_case_leaves_all_not_run(self):
        with patch.object(m, "check_device", return_value=False), patch.object(m, "append_test_run_history"), patch.object(m.socketio, "emit"):
            m.run_test_steps()
        self.assertEqual(m._last_workbook_run.statuses, ("not_run",) * len(self.steps))


if __name__ == "__main__":
    unittest.main()

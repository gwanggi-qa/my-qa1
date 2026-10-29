import io
import json
import threading
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from test_tc_drafts import DraftTests, result
from ai_request_errors import AIRequestError
import draft_batches as b
import tc_drafts as d


class BatchTests(DraftTests):
    # Inherited legacy tests also ensure historical generation remains readable.
    def setup_coverage(self, count=2):
        self.settings['count'] = 0
        self.calls = []
        def runner(instructions, context, metrics, schema, cancel, stop_type, image):
            self.calls.append(context.get('coverage_unit', {}).get('title', 'plan'))
            if 'units' in schema['properties']:
                return {'units':[{'title':str(i), 'source_ids':[context['evidence'][0]['id']]} for i in range(count)]}, {}
            value = result(context)
            value['cases'][0]['title'] = context['coverage_unit']['title']
            return value, {'input_tokens': 3}
        self.service.runner = runner

    def test_chunks_without_review_and_eta(self):
        self.setup_coverage()
        job = self.generate()
        self.assertEqual(job['status'], 'completed')
        self.assertIsNone(job['timeout_seconds'])
        self.assertEqual(job['progress']['completed'], 2)
        self.assertEqual(len(job['cases']), 2)
        book = load_workbook(io.BytesIO(self.service.export(self.sha, job['id'])))
        self.assertEqual(book['TestCase'].max_row, 3)
        self.assertIn('미검토', book['TestCase']['N2'].value)
        book.close()

    def test_stop_late_response_resume_preserves_first_chunk(self):
        self.setup_coverage()
        base = self.service.runner
        entered, release = threading.Event(), threading.Event()
        def slow(*args):
            if args[1].get('coverage_unit', {}).get('title') == '1':
                entered.set()
                release.wait(5)
            return base(*args)
        self.service.runner = slow
        job = self.service.start(self.sha, self.request())
        self.assertTrue(entered.wait(5))
        partial = self.service.get(self.sha, job['id'])
        self.assertEqual(len(partial['cases']), 1)
        self.assertGreaterEqual(partial['progress']['remaining_seconds'], 0)
        self.service.stop(self.sha, job['id'])
        release.set()
        self.service.worker.join(5)
        self.assertEqual(len(self.service.get(self.sha, job['id'])['cases']), 1)
        self.service.runner = base
        b.resume(self.service, self.sha, job['id'])
        self.service.worker.join(5)
        final = self.service.get(self.sha, job['id'])
        self.assertEqual(final['status'], 'completed')
        self.assertEqual([c['id'] for c in final['cases']], ['DTC-001','DTC-002'])
        self.assertEqual(self.calls.count('0'), 1)

    def test_failure_preserves_checkpoint_and_manual_resume(self):
        self.setup_coverage()
        base = self.service.runner
        def broken(*args):
            if args[1].get('coverage_unit', {}).get('title') == '1':
                raise AIRequestError('authentication')
            return base(*args)
        self.service.runner = broken
        job = self.generate()
        self.assertEqual(job['status'], 'failed')
        self.assertEqual(job['error_code'], 'authentication')
        self.assertEqual(len(job['cases']), 1)
        # A fresh service reads the persisted plan/cases; no original refs needed.
        self.service = d.DraftService(self.store, self.refs, base, self.status)
        b.resume(self.service, self.sha, job['id'])
        self.service.worker.join(5)
        self.assertEqual(self.service.get(self.sha, job['id'])['progress']['completed'], 2)

    def test_no_thousand_unit_cap(self):
        self.setup_coverage(1001)
        base = self.service.runner
        def halt(*args):
            if 'coverage_unit' in args[1]:
                raise d.DraftStopped()
            return base(*args)
        self.service.runner = halt
        job = self.generate()
        self.assertEqual(job['progress']['total'], 1001)
        self.assertTrue(job['progress']['resumable'])

    def test_invalid_evidence_is_not_saved_as_plan(self):
        self.setup_coverage()
        self.service.runner = lambda *args: ({'units':[{'title':'bad','source_ids':['missing']}]}, {})
        job = self.generate()
        self.assertEqual(job['error_code'], 'invalid_response')
        self.assertEqual(job['cases'], [])

    def test_total_duration_can_exceed_ten_minutes(self):
        self.setup_coverage()
        ticks = iter([0, 10, 410, 420, 820, 830])
        with patch.object(b.time, 'monotonic', side_effect=lambda: next(ticks)):
            job = self.generate()
        self.assertEqual(job['status'], 'completed')
        self.assertGreater(job['elapsed_seconds'], 600)

    def test_rejected_case_remains_excluded_from_export(self):
        self.setup_coverage()
        job = self.generate()
        case = job['cases'][0]
        self.service.review(self.sha, job['id'], case['id'],
                            {**{k:case[k] for k in d.FIELDS},'review_status':'rejected','revision':0})
        book = load_workbook(io.BytesIO(self.service.export(self.sha, job['id'])))
        self.assertEqual(book['TestCase'].max_row, 2)
        self.assertEqual(book['TestCase']['A2'].value, 'DTC-002')
        book.close()


if __name__ == '__main__':
    unittest.main()

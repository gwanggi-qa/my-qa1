import hashlib
import io
import json
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

import ai_tc
import mirror_server as m
import narrative_excel as ne


HEADERS = ['통합ID', '제목', '우선순위', '1Depth', 'Type', '테스트목적',
           '사전조건', '테스트절차', '기대결과', '자동화가능', 'Enabled']
PURPOSE = 'GNB 학습 관리 하위가 실력 분석 / 학습 현황 / 학습 로드맵 3종확인'


def workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = '학습_관리'
    ws.append(HEADERS)
    ws.append(['MGT-001', '학습 관리', 'P2', '학습_관리', 'UI-UX', PURPOSE,
               '학습 관리 화면 진입', '1. 메뉴 확인\n2. 스크롤 확인',
               '3종 모두 존재(사이드바 스크롤 시 확인)', 'FALSE', 'Y'])
    return wb


def content(wb):
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def parse(wb):
    return m.parse_excel(FileStorage(stream=content(wb), filename='qa.xlsx'))


def hint():
    return {'index': 1, 'intent': '메뉴 확인', 'actions': [], 'preconditions': [],
            'checks': [{'source': PURPOSE, 'text': s, 'present': True}
                       for s in ['실력 분석', '학습 현황', '학습 로드맵']], 'unsupported_reason': ''}


class NarrativeExcelTests(unittest.TestCase):
    def setUp(self):
        self.old_steps, self.old_metadata = m._test_steps, m._test_metadata
        m._test_steps, m._test_metadata = [{'action': 'wait', 'wait_ms': 10}], {'filename': 'previous.xlsx'}
        m._test_stop_event.clear()
        self.addCleanup(self.restore)
        self.client = m.app.test_client()

    def restore(self):
        m._test_steps, m._test_metadata = self.old_steps, self.old_metadata
        m._test_stop_event.clear()

    def post(self, wb, **fields):
        return self.client.post('/api/testcase', data={'file': (content(wb), 'qa.xlsx'), **fields})

    def test_false_is_metadata_type_is_not_action_and_blank_formatting_ignored(self):
        wb = workbook()
        wb.active.cell(920, 26).number_format = '0'
        steps = parse(wb)
        self.assertEqual(len(steps), 1)
        s = steps[0]
        self.assertEqual((s['case_no'], s['check_type'], s['source']), ('MGT-001', 'UI-UX', 'narrative_qa'))
        self.assertEqual(s['automation_possible'], 'FALSE')
        self.assertIn('\n', s['test_item'])
        self.assertEqual(s['required_labels'], ['실력 분석', '학습 현황', '학습 로드맵'])
        self.assertEqual(ai_tc.hard_block_reason(m, s), '')
        self.assertEqual(m.compact_test_step(1, s)['test_purpose'], PURPOSE)

    def test_actual_toeic_file(self):
        path = Path('C:/Users/USER/Desktop/뇌새김 TOEIC학습관리.xlsx')
        if not path.exists():
            self.skipTest('Local acceptance workbook unavailable')
        with path.open('rb') as f:
            steps = m.parse_excel(FileStorage(stream=f, filename=path.name))
        self.assertEqual([s['case_no'] for s in steps], [f'MGT-{i:03}' for i in range(1, 12)])
        self.assertTrue(all(s['visual_verification'] and not s['verification_gap'] for s in steps[1:4]))
        self.assertTrue(all('원본 학습 데이터' in s['verification_gap'] for s in steps[4:]))

    def test_explicit_disabled_and_merged_cells(self):
        wb = workbook()
        ws = wb.active
        ws.append(['MGT-002', 'Second', '', '', 'UI-UX', '', None, 'click', None, 'FALSE', 'Y'])
        ws.append(['MGT-003', 'Third', '', '', 'UI-UX', '', None, 'click', None, 'TRUE', 'N'])
        ws.merge_cells('G2:G3')
        steps = parse(wb)
        self.assertEqual(len(steps), 2)
        self.assertEqual(steps[1]['precondition'], steps[0]['precondition'])
        self.assertEqual(steps[1]['expected'], '')
        self.assertIn('기대결과', steps[1]['verification_gap'])

    def test_merged_context_does_not_create_cases_from_empty_rows(self):
        wb = workbook()
        wb.active.merge_cells('F2:F8')
        self.assertEqual(len(parse(wb)), 1)

    def test_title_rows_and_unknown_headers_mapping_preview(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Custom'
        ws.append(['QA document'])
        ws.append(['IdentifierX', 'BeforeX', 'DoX', 'OutcomeX'])
        ws.append(['X-1', 'Home', 'Click Next', 'Result\nScore'])
        response = self.post(wb)
        self.assertTrue(response.json['mapping_required'])
        self.assertEqual(m._test_metadata['filename'], 'previous.xlsx')
        selection = {'sheet': 'Custom', 'header_row': 2,
                     'mapping': {'case_no': 0, 'precondition': 1, 'test_item': 2, 'expected': 3}}
        response = self.post(wb, mapping=json.dumps(selection), preview='true')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['steps'][0]['expected'], 'Result\nScore')
        self.assertEqual(m._test_metadata['filename'], 'previous.xlsx')
        response = self.post(wb, mapping=json.dumps(selection))
        self.assertEqual(response.json['source'], 'narrative_qa')
        self.assertEqual(m._test_steps[0]['case_no'], 'X-1')

    def test_multiple_sheets_and_duplicate_columns_need_confirmation(self):
        wb = workbook()
        wb.copy_worksheet(wb.active)
        self.assertTrue(self.post(wb).json['mapping_required'])
        wb = workbook()
        wb.active.cell(1, 12, '기대결과')
        self.assertTrue(self.post(wb).json['mapping_required'])

    def test_invalid_mapping_and_upload_do_not_erase_previous(self):
        selection = {'sheet': '학습_관리', 'header_row': 1, 'mapping': {'test_item': 7, 'expected': 7}}
        self.assertEqual(self.post(workbook(), mapping=json.dumps(selection)).status_code, 400)
        self.assertEqual(m._test_metadata['filename'], 'previous.xlsx')
        response = self.client.post('/api/testcase', data={'file': (io.BytesIO(b'broken'), 'qa.xlsx')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(m._test_metadata['filename'], 'previous.xlsx')

    def test_required_list_cannot_be_reduced_to_one_label(self):
        step = parse(workbook())[0]
        ai_tc.validate_hint(hint(), step)
        partial = hint()
        partial['checks'].pop()
        with self.assertRaisesRegex(ValueError, 'mandatory'):
            ai_tc.validate_hint(partial, step)
        with patch.object(m, 'dump_ui_xml', return_value='<hierarchy><node text="실력 분석"/></hierarchy>'), \
                patch.object(m, 'ocr_screen_texts', return_value=[]):
            with self.assertRaises(AssertionError):
                ai_tc.verify(m, hint()['checks'])

    def test_narrative_batch_receives_live_candidates_and_image(self):
        steps = parse(workbook())
        screen = [{'candidate_id': 'c001', 'label': '学習', 'action': 'tap'}]
        with patch.object(m, 'ai_test_screen_candidates', return_value=([], screen)), \
                patch.object(m, 'capture_screen_png', return_value=b'png'), \
                patch.object(m, 'request_test_ai_json', return_value=({'hints': [hint()]}, {})) as ask:
            plans, _ = m.request_ai_test_batch_hints(steps, 1, {}, {})
        self.assertEqual(ask.call_args.args[1]['screen_candidates'], screen)
        self.assertEqual(ask.call_args.kwargs['image_png'], b'png')
        self.assertEqual(len(plans[1]['checks']), 3)

    def test_known_gaps_do_not_consume_ai_requests_or_renumber_cases(self):
        steps = parse(workbook())
        blocked = {**steps[0], 'case_no': 'DATA', 'verification_gap': 'Needs source data'}
        reply = hint()
        reply['index'] = 2
        with patch.object(m, 'ai_test_screen_candidates', return_value=([], [])), \
                patch.object(m, 'capture_screen_png', return_value=None), \
                patch.object(m, 'request_test_ai_json', return_value=({'hints': [reply]}, {})) as ask:
            plans, _ = m.request_ai_test_batch_hints([blocked, steps[0], blocked], 1, {}, {})
            self.assertEqual(list(plans), [1, 3, 2])
            self.assertEqual([s['index'] for s in ask.call_args.args[1]['steps']], [2])
            ask.reset_mock()
            plans, _ = m.request_ai_test_batch_hints([blocked], 11, {}, {})
            ask.assert_not_called()
            self.assertEqual(plans[11]['unsupported_reason'], 'Needs source data')

    def test_explicit_scroll_can_reveal_labels_but_cannot_navigate_away(self):
        tc = {**parse(workbook())[0], '_ai_phase': 'verify', 'ai_hint': hint()}
        snapshot = hashlib.sha256(json.dumps([], sort_keys=True).encode()).hexdigest()
        candidate = {'action': 'swipe', 'source': 'ui-analysis', 'label': 'Sidebar', 'bounds': '[0,0][200,700]'}
        recovery = {'decision': 'execute', 'action': 'swipe', 'candidate_id': 'c1', 'candidate': candidate,
                    'reason': 'Reveal lower menu', 'wait_ms': 0, 'snapshot': snapshot}
        with patch.object(m, 'request_ai_test_recovery', return_value=recovery), \
                patch.object(m, 'ai_test_screen_candidates', return_value=([], [])), \
                patch.object(m, 'capture_failure_screenshot', return_value='proof.png'), \
                patch.object(m, 'emit_test_ai_status'), \
                patch.object(ai_tc, 'execute_action', return_value='swiped') as execute, \
                patch.object(ai_tc, 'execute_plan', return_value='All three labels verified'):
            result = m.run_ai_test_recovery(1, tc, 'Missing lower menu', {}, {})
            self.assertIn('verified', result['message'])
            execute.assert_called_once()
            recovery['action'] = 'tap'
            with self.assertRaisesRegex(m.AIRecoveryExhausted, 'only an observed scrollable'):
                m.run_ai_test_recovery(1, tc, 'Missing lower menu', {}, {})
            self.assertEqual(execute.call_count, 1)

    def test_expected_scroll_clause_can_ground_swipe(self):
        tc = parse(workbook())[0]
        value = hint()
        value['actions'] = [{'action': 'swipe', 'source': tc['expected'], 'target': 'Sidebar',
                             'direction': 'up', 'wait_ms': 0, 'value': ''}]
        ai_tc.validate_hint(value, tc)

    def test_data_needs_evidence_visual_uses_screenshot_verifier(self):
        for expected in ['환산 점수가 정확하게 반올림된다', '카드 안에 LC/RC 토글이 노출된다']:
            wb = workbook()
            wb.active.cell(2, 6, '')
            wb.active.cell(2, 9, expected)
            step = parse(wb)[0]
            step['ai_hint'] = hint()
            with patch.object(ai_tc, 'execute_action') as execute, \
                    patch.object(m, 'verify_visual_step', side_effect=m.BlockedStep('No visual evidence')) as visual:
                with self.assertRaises(m.BlockedStep):
                    ai_tc.execute_plan(m, step)
                self.assertEqual(visual.call_count, int(step['visual_verification']))
            execute.assert_not_called()

    def test_standard_run_blocks_pending_and_missing_ai_plan(self):
        step = parse(workbook())[0]
        with self.assertRaisesRegex(m.BlockedStep, 'AI'):
            m.execute_step(step)
        with self.assertRaisesRegex(m.BlockedStep, 'AI plan'):
            ai_tc.execute_plan(m, step)

    def test_literal_text_check_works_without_ai(self):
        wb = Workbook()
        wb.active.append(['통합ID', '테스트목적', '테스트절차', '기대결과'])
        wb.active.append(['TXT-1', '', '', '"완료" 표시'])
        step = parse(wb)[0]
        self.assertEqual(step['action'], 'assert_text')
        self.assertFalse(step['ai_pending'])
        with patch.object(m, 'wait_until_any_ui_contains', return_value='완료') as observe:
            self.assertEqual(m.execute_step(step), 'Found text: 완료')
        self.assertIn('완료', observe.call_args.args[0])

    def test_recorder_and_manual_qa_keep_their_format(self):
        wb = Workbook()
        wb.active.append(['index', 'type', 'x', 'y', 'delay_ms', 'screen_width', 'screen_height'])
        wb.active.append([1, 'tap', 10, 20, 0, 800, 1340])
        self.assertEqual(parse(wb)[0]['source'], 'recorder')
        wb = Workbook()
        wb.active.append(['1Depth', 'Check', 'Expected Result'])
        wb.active.append(['홈', 'UI 리소스', '"시작하기" 표시'])
        self.assertEqual(parse(wb)[0]['source'], 'manual_qa')


if __name__ == '__main__':
    unittest.main()

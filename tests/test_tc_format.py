import io
import unittest
import zipfile

from openpyxl import load_workbook
import result_workbook as r
import tc_format as f


class StandardFormatTests(unittest.TestCase):
    def test_verdicts_preserve_manual_expected_and_evidence(self):
        steps=[{'case_no':f'TC-{i}', 'sheet':'TestCase','row':i+1,
                'expected':'Expected original','qa_result':'manual keep','note':'=SUM(1,2)'}
               for i in range(1,7)]
        book=f.workbook(f.step_row(s) for s in steps)
        book.create_sheet('Sources').append(['Original evidence'])
        raw=f.save(book)
        template=r.prepare_template(raw,'testcase.xlsx',steps)
        statuses=['passed','failed','blocked','stopped','not_run','not_applicable']
        snapshot=r.finish_run('a'*32,template,steps,
                              [{'index':i,'status':s} for i,s in enumerate(statuses,1)])
        output=r.export_run(snapshot,template.file_hash)
        run={'run_id':'a'*32,'results':[{'source':{'sheet':'TestCase','row':2,'tc_id':'TC-1','expected':'Expected original'},
             'status':'passed','actual':'=HYPERLINK("untrusted")','reason':'','evidence':[{'url':'/fixture.png'}],'stages':[]}]}
        output=f.append_execution_details(output,run)
        with zipfile.ZipFile(io.BytesIO(raw)) as original,zipfile.ZipFile(io.BytesIO(output)) as saved:
            self.assertEqual(original.read('xl/worksheets/sheet2.xml'),saved.read('xl/worksheets/sheet2.xml'))
        check=load_workbook(io.BytesIO(output))
        self.assertEqual([c.value for c in check['TestCase'][1]],list(f.HEADERS))
        for row,status in enumerate(statuses,2):
            self.assertEqual(check['TestCase'].cell(row,11).value,'Expected original')
            self.assertEqual(check['TestCase'].cell(row,12).value,'manual keep')
            self.assertEqual(check['TestCase'].cell(row,13).value,f.STATUS_CODES[status])
            self.assertEqual(check['TestCase'].cell(row,14).data_type,'s')
        details=check[check.sheetnames[-1]]
        self.assertEqual(details['F2'].data_type,'s')
        self.assertEqual(details['G2'].value,'/fixture.png')
        self.assertTrue(details['F2'].alignment.wrap_text)
        check.close()

    def test_blank_template_and_repeated_detail_name(self):
        book=f.workbook()
        self.assertEqual(book.active.max_row,1)
        self.assertEqual(book.active.freeze_panes,'A2')
        raw=f.save(book)
        run={'run_id':'b'*32,'results':[]}
        for _ in range(2):
            raw=f.append_execution_details(raw,run)
        check=load_workbook(io.BytesIO(raw))
        self.assertEqual(len(set(check.sheetnames)),3)
        self.assertTrue(all(len(n)<=31 for n in check.sheetnames))
        check.close()

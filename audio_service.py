"""Subject-specific source-audio checks; no device input or issue publication."""
import copy
import hashlib
import io
import json
import re
import threading
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin, urlsplit, parse_qs, urlencode
import requests
from bs4 import BeautifulSoup
from flask import jsonify, request, send_file, make_response
from audio_verification import classify

BASE = 'https://theweek.brain-study.co.kr/'
ROOT = Path(__file__).parent / 'qa_results' / 'audio_verification'
DEFAULT_SUBJECT = 'japanese'
SUBJECTS = {
    'toeic': dict(name='뇌새김 TOEIC 완성', prefix='toe', app='toe_us', apps=['toe_us','toe_uk','toe_au','toe_ca'], language='en', unit_label='Unit', level_names={'1':'빈출편 1권','2':'빈출편 2권','3':'빈출편 3권','4':'빈출편 4권','5':'착각편'}),
    'travel-english': dict(name='더위크 여행영어', prefix='t', app='trip_c', apps=['trip_w','trip_c'], language='en', unit_label='단원', level_names={'1':'스타터','2':'부스터','3':'마스터','4':'핸디북'}),
    'chinese': dict(name='더위크 중국어', prefix='c', app='cn', language='zh', unit_label='Day'),
    'japanese': dict(name='더위크 일본어', prefix='j', app='jp', language='ja', unit_label='Day'),
    'japanese-premium': dict(name='더위크 일본어 프리미엄', prefix='tjp2', app='tjp2', language='ja', unit_label='Stage / Session'),
}
LABELS = {'match': '일치', 'review': '검토 필요', 'mismatch_candidate': '불일치 후보', 'unavailable': '검사 불가'}
STAGES = {'word': '단어 학습', 'sentence': '문장 학습', 'conversation': '회화 학습', 'other': '기타'}
REASONS = {
    'Short utterance requires listening review': '음성이 짧아 정확한 인식이 어려울 수 있습니다. 직접 듣고 확인해 주세요.',
    'Spelling or speech recognition ambiguity': '표기 차이인지 음성 인식 오류인지 확실하지 않습니다. 직접 듣고 확인해 주세요.',
    'Normalized transcript matches display text': '숫자 표기와 문장부호 차이를 제외하면 화면 문구와 일치합니다.',
    'Numeric content differs; confirm by listening': '화면과 인식된 음성의 숫자가 다릅니다. 직접 듣고 확인해 주세요.',
    'Question/statement differs; confirm by listening': '화면 문구와 음성의 질문·설명 형식이 다릅니다. 직접 듣고 확인해 주세요.',
    'Missing text, transcript or audio': '화면 문구, 음성 인식 결과 또는 음원이 없어 검사할 수 없습니다.',
    'Audio download, decoding or transcription failed': '음원 다운로드 또는 음성 인식에 실패했습니다.',
}
UNIT_STATES = {'finished': '검사 완료', 'stopped': '중지됨', 'failed': '검사 실패', 'unavailable': '검사 불가'}

def now():
    return datetime.now(timezone.utc).isoformat()

def fetch(url, limit):
    p = urlsplit(url)
    allowed = {'/seoch.php'} | {f'/{s["prefix"]}{page}.php' for s in SUBJECTS.values() for page in ('level','week','day')}
    if p.scheme != 'https' or p.netloc != 'theweek.brain-study.co.kr' or p.path not in allowed:
        raise ValueError('unsupported_source')
    with requests.get(url, timeout=(10, 30), stream=True, allow_redirects=False) as r:
        if r.status_code != 200:
            raise ValueError('source_http_error')
        chunks, size = [], 0
        for chunk in r.iter_content(65536):
            size += len(chunk)
            if size > limit:
                raise ValueError('source_too_large')
            chunks.append(chunk)
        return b''.join(chunks)

def discover_catalog(subject=DEFAULT_SUBJECT):
    prefix = SUBJECTS[subject]['prefix']
    soup = BeautifulSoup(fetch(BASE+prefix+'level.php', 2000000), 'html.parser')
    levels = set()
    for node in soup.select('[onclick]'):
        m = re.fullmatch(r'func\.goWeek\((\d+)\);?', node['onclick'].strip())
        if m:
            levels.add(int(m[1]))
    if not levels or len(levels) > 30:
        raise ValueError('invalid_catalog')
    units = {}
    for level in sorted(levels):
        page = BeautifulSoup(fetch(BASE+f'{prefix}week.php?l={level}', 2000000), 'html.parser')
        for node in page.select('[onclick], a[href]'):
            value = node.get('href', '')
            if node.has_attr('onclick'):
                m = re.fullmatch(r"(?:window\.)?location\.href\s*=\s*['\"]([^'\"]+)['\"];?", node['onclick'].strip())
                value = m[1] if m else ''
            p = urlsplit(urljoin(BASE,value)); q = parse_qs(p.query)
            if p.netloc != urlsplit(BASE).netloc or p.path != f'/{prefix}day.php':
                continue
            if subject == 'toeic' and 'w' not in q:
                q['w']=['1']
            if any(len(q.get(k,[])) != 1 or not q[k][0].isdigit() for k in ('l','w','d')):
                continue
            l,w,d = (int(q[k][0]) for k in ('l','w','d'))
            if l != level or not all(0 < v < 1000 for v in (l,w,d)):
                continue
            key = f'L{l}-W{w}-D{d}'
            params=dict(l=l,w=w,d=d); session=None
            if subject == 'japanese-premium':
                if len(q.get('s',[])) != 1 or q['s'][0] not in ('1','2','3'): continue
                session=int(q['s'][0]); params['s']=session; key+=f'-S{session}'
            units[key] = dict(key=key,subject=subject,level=l,week=w,day=d,session=session,
                             label=f'Level {l} / Week {w} / '+(f'Stage {d} / Session {session}' if session else f'Day {d}'),
                             url=BASE+prefix+'day.php?'+urlencode(params))
            if subject == 'travel-english':
                title=node.get_text(' ',strip=True)
                units[key]['label']=f"{SUBJECTS[subject]['level_names'][str(l)]} / 그룹 {w} / {title or '단원 '+str(d)}"
                units[key]['title']=title
            if subject == 'toeic':
                units[key]['url']=BASE+prefix+'day.php?'+urlencode(dict(l=l,d=d))
                units[key]['label']=SUBJECTS[subject]['level_names'].get(str(l),str(l))+f' / Unit {d}'
    if not units or len(units)>1000:
        raise ValueError('invalid_catalog')
    return sorted(units.values(),key=lambda u:(u['level'],u['week'],u['day'],u.get('session') or 0))

def parse_items(html,url,subject=None):
    soup=BeautifulSoup(html,'html.parser'); items=[]; stage='other'
    if subject == 'toeic':
        for node in list(soup.select('.playlist-item')):
            for accent,label in [('us','미국'),('uk','영국'),('au','호주'),('ca','캐나다')]:
                clone=copy.deepcopy(node)
                clone['audio_url']=node.get('data-url-'+accent,'')
                clone['data-title']=label+' 발음'
                node.insert_before(clone)
            node.decompose()
    for node in soup.select('.playlist_text, .playlist-item'):
        if 'playlist_text' in node.get('class',[]):
            title=node.get_text(' ',strip=True)
            stage=next((v for k,v in [('단어','word'),('문장','sentence'),('회화','conversation')] if k in title),title)
            continue
        title=node.select_one('.playlist-subtitle')
        source=urljoin(url,node['audio_url']) if node.get('audio_url') else ''
        if subject and source:
            p=urlsplit(source)
            apps=SUBJECTS[subject].get('apps',[SUBJECTS[subject]['app']])
            app_values=parse_qs(p.query).get('app',[])
            if p.netloc != urlsplit(BASE).netloc or p.path != '/seoch.php' or len(app_values)!=1 or app_values[0] not in apps:
                raise ValueError('subject_audio_mismatch')
        items.append(dict(index=len(items)+1,stage=node.get('data-title') or stage,expected=title.get_text(' ',strip=True) if title else '',
                          source_url=urljoin(url,node['audio_url']) if node.get('audio_url') else '',
                          status='pending',review='pending',item_capture=None))
    if not 1<=len(items)<=(400 if subject == 'toeic' else 100):
        raise ValueError('invalid_item_count')
    return items

def load_model():
    from faster_whisper import WhisperModel
    return WhisperModel(str(Path(__file__).parent/'qa_results'/'fw-small'),device='cpu',compute_type='int8')

def atomic_json(path,data):
    temp=path.with_suffix('.tmp')
    temp.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    temp.replace(path)

def public(data):
    out=copy.deepcopy(data)
    out.setdefault('subject',DEFAULT_SUBJECT)
    out['subject_name']=SUBJECTS[out['subject']]['name']
    for u in out.get('units',[]):
        u.pop('html',None)
    return out

def register(app,server,root=None):
    root=Path(root or ROOT); root.mkdir(parents=True,exist_ok=True)
    lock=threading.RLock(); stop=threading.Event(); state={'current':None}; catalog={}; previews={}

    def paths():
        return [p for p in list(root.glob('*/result.json')) + list(root.glob('*/*/result.json'))
                if re.fullmatch(r'[a-f0-9]{32}',p.parent.name)]

    def folder_for(run_id):
        if not re.fullmatch(r'[a-f0-9]{32}',run_id): return root/'invalid'
        for parent in [root]+[root/s for s in SUBJECTS]:
            if (parent/run_id/'result.json').is_file(): return parent/run_id
        return root/'invalid'

    def get_catalog(subject=DEFAULT_SUBJECT):
        if subject not in catalog:
            catalog[subject]=discover_catalog(subject)
        return catalog[subject]

    def save(data):
        folder=root/data['subject']/data['id']
        data['items']=[dict(x,subject=data['subject'],unit=u['key'],level=u['level'],week=u['week'],day=u['day'],session=u.get('session')) for u in data['units'] for x in u['items'] if x['status']!='pending']
        data['total']=sum(len(u['items']) for u in data['units'])
        data['counts']=dict(Counter(x['status'] for x in data['items']))
        for u in data['units']:
            u['counts']=dict(Counter(x['status'] for x in u['items'] if x['status']!='pending'))
            atomic_json(folder/u['key']/'result.json',{k:v for k,v in u.items() if k!='html'})
        atomic_json(folder/'result.json',public(data))

    def run(data):
        folder=root/data['subject']/data['id']
        try:
            try:
                (folder/'start.png').write_bytes(server().capture_screen_png()); data['capture']='start.png'
            except Exception:
                data['capture']=None
            model=load_model()
            for u in data['units']:
                if stop.is_set(): break
                if u['status'] in ('finished','unavailable'): continue
                u['status']='running'
                for item in u['items']:
                    if stop.is_set(): break
                    if item['status']!='pending': continue
                    with lock:
                        data.update(active_unit=u['key'],active_index=item['index']); save(data)
                    result=copy.deepcopy(item)
                    try:
                        if not result['expected'] or not result['source_url']: raise ValueError('missing_source')
                        raw=fetch(result['source_url'],10000000)
                        name=f"{item['index']:02d}.mp3"; path=folder/u['key']/name; path.write_bytes(raw)
                        result.update(audio=f"{u['key']}/{name}",sha256=hashlib.sha256(raw).hexdigest())
                        segments,info=model.transcribe(str(path),language=SUBJECTS[data['subject']]['language'],beam_size=5,condition_on_previous_text=False)
                        transcript=''.join(s.text for s in segments).strip()
                        if stop.is_set(): break
                        result.update(transcript=transcript,duration=info.duration)
                        result.update(classify(result['expected'],transcript,duration=info.duration,language=SUBJECTS[data['subject']]['language']))
                        path.with_suffix('.txt').write_text(transcript,encoding='utf-8')
                    except Exception:
                        result.update(status='unavailable',reason='Audio download, decoding or transcription failed')
                    with lock:
                        item.update(result); save(data)
                    if stop.wait(.1): break
                with lock:
                    u['status']='stopped' if stop.is_set() else 'finished'; save(data)
            outcome='stopped' if stop.is_set() else 'finished'
        except Exception:
            outcome='failed'
        finally:
            with lock:
                data.update(status=outcome,finished_at=now())
                if outcome=='failed': data['error']='Local speech model or result storage unavailable'
                save(data)

    def start(units,resumed_from=None,subject=DEFAULT_SUBJECT):
        data=dict(id=uuid.uuid4().hex,subject=subject,subject_name=SUBJECTS[subject]['name'],status='running',units=copy.deepcopy(units),items=[],started_at=now(),
                  method='server-source-audio / local Whisper small',capture_scope='run-start-only',redmine='not_connected',resumed_from=resumed_from)
        folder=root/subject/data['id']; folder.mkdir(parents=True)
        for u in data['units']:
            dest=folder/u['key']; dest.mkdir()
            if resumed_from:
                import shutil
                for old in (folder_for(resumed_from)/u['key']).iterdir():
                    if old.is_file() and re.fullmatch(r'(source\.html|\d{2,3}\.(mp3|txt))',old.name): shutil.copy2(old,dest/old.name)
                if u['status'] not in ('finished','unavailable'): u['status']='pending'
            elif u.get('html'):
                (dest/'source.html').write_text(u['html'],encoding='utf-8')
            u.pop('html',None)
        state['current']=data; stop.clear(); save(data)
        threading.Thread(target=run,args=(data,),daemon=True).start()
        return jsonify(id=data['id']),202

    for path in sorted(paths(),key=lambda p:p.stat().st_mtime,reverse=True):
        try:
            old=json.loads(path.read_text(encoding='utf-8'))
            if old.get('status')=='running':
                old.update(status='stopped',error='Server restarted',finished_at=now()); atomic_json(path,old)
            state['current']=old; break
        except (ValueError,OSError): continue

    @app.get('/audio-verification')
    def audio_page():
        html=(Path(__file__).parent/'static'/'audio-verification.html').read_text(encoding='utf-8')
        return server().local_access.bootstrap(make_response(html.replace('__TOKEN__',server().local_access.token)))

    @app.get('/api/audio-verification/catalog')
    def audio_catalog():
        subject=request.args.get('subject',DEFAULT_SUBJECT)
        if subject not in SUBJECTS: return jsonify(error='지원하지 않는 과목입니다.'),400
        try: return jsonify(subject=subject,units=get_catalog(subject))
        except Exception: return jsonify(error='단원 목록을 가져오지 못했습니다. 연결 상태를 확인해 주세요.'),502

    @app.get('/api/audio-verification/subjects')
    def audio_subjects():
        return jsonify(subjects=[dict(id=k,**v) for k,v in SUBJECTS.items()])

    @app.post('/api/audio-verification/preview')
    def audio_preview():
        body=request.get_json(silent=True) or {}; keys=body.get('units') if isinstance(body,dict) else None
        subject=body.get('subject',DEFAULT_SUBJECT) if isinstance(body,dict) else None
        if subject not in SUBJECTS: return jsonify(error='지원하지 않는 과목입니다.'),400
        if not isinstance(keys,list) or not keys or len(keys)>200 or any(not isinstance(k,str) for k in keys) or len(keys)!=len(set(keys)):
            return jsonify(error='검사할 단원을 1~200개 선택해 주세요.'),400
        try: known={u['key']:u for u in get_catalog(subject)}
        except Exception: return jsonify(error='단원 목록 연결 오류'),502
        if any(k not in known for k in keys): return jsonify(error='실제 목록에 없는 단원입니다.'),400
        units=[]
        for k in keys:
            u=copy.deepcopy(known[k]); u.update(subject=subject,status='pending',items=[])
            try:
                u['html']=fetch(u['url'],2000000).decode('utf-8'); u['items']=parse_items(u['html'],u['url'],subject)
            except Exception: u.update(status='unavailable',error='단원 문구 또는 음원 목록을 읽지 못했습니다.')
            units.append(u)
        units.sort(key=lambda u:(u['level'],u['week'],u['day'])); pid=uuid.uuid4().hex
        with lock:
            while len(previews)>=4: previews.pop(next(iter(previews)))
            previews[pid]=(subject,units)
        return jsonify(subject=subject,preview_id=pid,units=public({'units':units})['units'],total=sum(len(u['items']) for u in units))

    @app.route('/api/audio-verification',methods=['GET','POST'])
    def audio_runs():
        with lock:
            if request.method=='GET': return jsonify(current=public(state['current']) if state['current'] else None)
            if state['current'] and state['current']['status']=='running': return jsonify(error='이미 음원 검사가 진행 중입니다.'),409
            body=request.get_json(silent=True) or {}
            pid=body.get('preview_id') if isinstance(body,dict) else None
            prepared=previews.get(pid) if isinstance(pid,str) else None
            subject,units=prepared if prepared else (DEFAULT_SUBJECT,None)
            if body.get('subject',subject) != subject: return jsonify(error='과목이 변경되었습니다. 미리보기를 다시 확인하세요.'),409
            if not units or not any(u['items'] for u in units): return jsonify(error='검사 대상 미리보기를 먼저 확인해 주세요.'),400
            return start(units,subject=subject)

    @app.post('/api/audio-verification/stop')
    def audio_stop():
        stop.set(); return jsonify(status='stopping')

    @app.post('/api/audio-verification/resume')
    def audio_resume():
        with lock:
            d=state['current']
            if not d or d['status'] not in ('stopped','failed') or not d.get('units') or not any(x['status']=='pending' for u in d['units'] for x in u['items']):
                return jsonify(error='재개 가능한 미완료 항목이 없습니다.'),409
            return start(d['units'],d['id'],d.get('subject',DEFAULT_SUBJECT))

    @app.get('/api/audio-verification/history')
    def audio_history():
        rows=[]
        subject=request.args.get('subject')
        if subject is not None and subject not in SUBJECTS: return jsonify(error='지원하지 않는 과목입니다.'),400
        for p in sorted(paths(),key=lambda p:p.stat().st_mtime,reverse=True):
            try:
                d=public(json.loads(p.read_text(encoding='utf-8')))
                if subject and d['subject']!=subject: continue
                rows.append({k:d.get(k) for k in ('id','subject','subject_name','started_at','status','counts')})
                if len(rows)>=100: break
            except (ValueError,OSError): continue
        return jsonify(runs=rows)

    @app.get('/api/audio-verification/<run_id>/export')
    def audio_export(run_id):
        folder=folder_for(run_id)
        if not (folder/'result.json').is_file(): return jsonify(error='not_found'),404
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        d=public(json.loads((folder/'result.json').read_text(encoding='utf-8'))); wb=Workbook(); ws=wb.active; ws.title='음원 검증'
        ws.append(['회차','Level','Week','Day','항목','학습 유형','화면 문구','음성에서 인식한 문장','판정','판정 이유','음원 파일','SHA-256'])
        for x in d.get('items',[]):
            ws.append([run_id,x.get('level'),x.get('week'),x.get('day'),x['index'],STAGES.get(x.get('stage'),x.get('stage')),x.get('expected'),x.get('transcript'),LABELS.get(x['status'],x['status']),REASONS.get(x.get('reason'),x.get('reason')),x.get('audio'),x.get('sha256')])
            for c in ws[ws.max_row]:
                if isinstance(c.value,str): c.data_type='s'
                c.alignment=Alignment(wrap_text=True,vertical='top')
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='25616B')
        for col in 'ABCDEFGHIJKL': ws.column_dimensions[col].width=32 if col in 'AGHJKL' else 14
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        summary=wb.create_sheet('단원 요약'); summary.append(['단원','상태','항목 수','일치','불일치 후보','검토 필요','검사 불가','오류'])
        for u in d.get('units',[]):
            c=u.get('counts',{}); summary.append([u['key'],UNIT_STATES.get(u['status'],u['status']),len(u['items']),c.get('match',0),c.get('mismatch_candidate',0),c.get('review',0),c.get('unavailable',0),u.get('error','')])
        ws.cell(1,13,'과목'); ws.cell(1,14,'Session'); ws.cell(1,15,'원본 음원 URL')
        for row,x in enumerate(d.get('items',[]),2):
            ws.cell(row,13,d['subject_name']); ws.cell(row,14,x.get('session')); ws.cell(row,15,x.get('source_url'))
        ws.column_dimensions['M'].width=30; ws.column_dimensions['O'].width=55
        ws.auto_filter.ref=ws.dimensions
        meta=wb.create_sheet('검사 정보'); meta.append(['과목',d['subject_name']]); meta.append(['과목 ID',d['subject']]); meta.append(['검사 방식',d.get('method')]); meta.append(['시작 시각',d.get('started_at')])
        out=io.BytesIO(); wb.save(out); out.seek(0)
        return send_file(out,as_attachment=True,download_name=f'audio-{run_id}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    @app.get('/api/audio-verification/<run_id>/<path:filename>')
    def audio_evidence(run_id,filename):
        if not re.fullmatch(r'[a-f0-9]{32}',run_id) or not re.fullmatch(r'(start\.png|(?:L\d+-W\d+-D\d+(?:-S\d+)?/)?(?:result\.json|\d{2,3}\.(?:mp3|txt)))',filename):
            return jsonify(error='not_found'),404
        path=folder_for(run_id)/filename
        if not path.is_file(): return jsonify(error='not_found'),404
        if filename=='result.json': return jsonify(public(json.loads(path.read_text(encoding='utf-8'))))
        return send_file(path)

    return state
